# -*- coding: utf-8 -*-
"""2026-09-14 — predmjer LOT 1 po komentarima recenzenta (A. Čolpa, 27.08.2026).

Komentari su dati na reviziju sa 3 nosača 2×2 portret; ovdje se primjenjuju na
usvojeni raspored. Zajednički predmjer (rfp-hybrid-solar-td, joint_boq.py) kopira
ovaj fajl, pa izmjene važe za obje lokacije.

  R1  1.1  „U stavku 1.1 bih dodao da je ponuđač obavezan da uz ponudu dostavi
           tehničko rješenje konstrukcije za prihvat fotonaponskih panela sa
           obaveznim statičkim proračunom i specifikacijom materijala za predmetnu
           lokaciju." → novi stav uz ponudu; projektne sile po nosaču više se ne
           zadaju brojem nego proizlaze iz tehničkog rješenja
  R2  1.2  „Treba izbjeći navođenje vrijednosti sila čupanja i podizanja, te
           momenata prevrtanja." → nosivost ankera i sile ankerisanja idu na ETA i
           statički proračun; uklonjeni ≥30 kN, ≥13,3 kN i ≥16,1 kN
  R3  1.3  „Predlažem da se izbace vrijednosti pritiska vjetra i koeficijenata
           sigurnosti... Ostaviti i naglasiti standarde." → ostaju samo standardi
           (dodat BAS EN 1993-1-1 za samu konstrukciju); napomena preformulisana
           njegovom rečenicom: proračun je uslov za POČINJANJE radova, ne za
           primopredaju
  R4  2.2  „Nije jasno o kakvom se zatrpavanju radi, ako je količina betona
           približno ili skoro identična iskopu." → rov se u stijeni siječe jednom
           širinom (500 mm), a traka ide punom dubinom, pa ga beton i podložni
           beton ispunjavaju u cijelosti: zatrpavanja nema i stavka se briše
  R5  2.x  „Prekontrolisati količinu iskopa / betona / stavku nakon usaglašavanja"
           → iskop ostaje 9,88 m³; odvoz 8,94 → 9,88 m³ (sav iskop); beton
           8,42 → 9,36 m³ (traka jedinstvene širine 500, cad/design.json); obrada
           12 → 10,4 m² (0,50 × 2,60 × 8). Sekcija 2 se prenumeriše 2.1…2.6 jer
           je „Zatrpavanje" ispalo; jedina unakrsna referenca (1.2 → Tačka 2.3)
           se povlači na 2.4.

Idempotentno kao fix_boq_sw_4x3l.py: drugi prolaz ništa ne mijenja i ne snima
fajl. Brisanje reda čini ranije jednokratne skripte (fix_boq_quantities.py,
fix_boq_sw_4x3l.py) neponovljivim — one su zapis primijenjene revizije, ne alat.

    python fix_boq_recenzija_a.py
"""
import os
import shutil
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fix_boq_dc_aux as dc  # noqa: E402  (XLSX, jb, xml_rows, props, assert_no_loss)

XLSX, jb = dc.XLSX, dc.jb

# ---------------------------------------------------------------- R1, R2, R3
EDITS = {
    "1.1": [
        ("Projektne sile po nosaču (GSN, γQ=1,5 / γG,fav=0,9): podizanje ≥13,3 kN, "
         "horizontalna sila ≥10,0 kN, moment prevrtanja ≥25,7 kNm. Mjerodavni su podizanje "
         "(uplift) i prevrtanje, a ne nosivost tla.",
         "Projektne sile po nosaču (podizanje, horizontalna sila, moment prevrtanja) "
         "proizlaze iz koncepta konstrukcije koji Ponuđač nudi; Ponuđač ih iskazuje i "
         "dokazuje u tehničkom rješenju i statičkom proračunu iz Tačke 1.3. Mjerodavni su "
         "podizanje (uplift) i prevrtanje, a ne nosivost tla."),
        ("po jedan string na svaku od dvije rute PVDB ormara",
         "po jedan string na svaku od dvije rute PVDB ormara\n"
         " - UZ PONUDU: Ponuđač je obavezan da uz ponudu dostavi tehničko rješenje "
         "konstrukcije za prihvat fotonaponskih panela za predmetnu lokaciju, sa obaveznim "
         "statičkim proračunom (Tačka 1.3) i specifikacijom materijala. Tehničko rješenje "
         "mora biti pregledano i odobreno prije početka izvođenja radova"),
    ],
    "1.2": [
        ("- nosivost: karakteristična sila čupanja ≥30 kN po ankeru; dubina ugradnje prema "
         "ETA za konkretnu podlogu. Ukupna projektna sila podizanja po nosaču ≥13,3 kN "
         "(GSN); sila po traci od momenta prevrtanja ≥16,1 kN pri razmaku traka 1600 mm",
         "- nosivost i dubina ugradnje: prema ETA za konkretnu podlogu i prema statičkom "
         "proračunu iz Tačke 1.3; projektne sile ankerisanja proizlaze iz koncepta "
         "konstrukcije koji Ponuđač nudi i iskazuju se u tehničkom rješenju"),
        ("dozvoljena su SAMO uz gravitacioni temelj iz Tačke 2.3 (dvije trake po nosaču, "
         "najmanje 0,74 m³ po traci; tendovano 1,053 m³ po traci) — momenat prevrtanja i "
         "dalje zahtijeva provjeru lokalnog istezanja trake, ne samo ukupne mase",
         "dozvoljena su SAMO uz gravitacioni temelj iz Tačke 2.4 (dvije trake po nosaču, "
         "tendovano 1,170 m³ po traci) — dovoljnost vlastite težine i lokalno istezanje "
         "trake dokazuju se statičkim proračunom iz Tačke 1.3"),
    ],
    "1.3": [
        ("- dokaz otpornosti na pritisak vjetra qp ≥ 1,20 kN/m² pri nagibu 45°, prema "
         "BAS EN 1991-1-4 sa BiH nacionalnim aneksom, uključujući orografiju",
         "- dokaz na dejstvo vjetra prema BAS EN 1991-1-4 sa BiH nacionalnim aneksom, "
         "uključujući faktor orografije za lokalitet"),
        ("- dokaz sigurnosti na podizanje (uplift) i prevrtanje prema EN 1990 "
         "(γQ,dst = 1,5; γG,stb = 0,9)",
         "- dokaz sigurnosti na podizanje (uplift) i prevrtanje prema BAS EN 1990\n"
         " - dimenzionisanje čelične konstrukcije prema BAS EN 1993-1-1, izvedba prema "
         "BAS EN 1090-2"),
        ("NAPOMENA: proračun se dostavlja tokom realizacije kao uslov za primopredaju.",
         "NAPOMENA: Statički proračun se dostavlja u sklopu tehničkog rješenja i uslov je "
         "za počinjanje radova. Prije početka izvođenja radova tehničko rješenje mora biti "
         "pregledano i odobreno."),
    ],
    # ------------------------------------------------------------ R4, R5
    "2.1": [
        ("4 nosača × 2 trake × 1,235 m3",
         "4 nosača × 2 trake × 1,235 m3. Rov se u stijeni siječe jednom širinom po cijeloj "
         "dubini; traka se betonira do zida rova, pa zatrpavanja nema."),
    ],
    "2.2a": [
        ("Odvoz i zbrinjavanje viška materijala od iskopa do deponije udaljene do 20 km.",
         "Odvoz i zbrinjavanje viška materijala od iskopa do deponije udaljene do 20 km. "
         "Traka i podložni beton ispunjavaju rov u cijelosti, pa se odvozi sav iskopani "
         "materijal. 4 nosača × 2 trake × 1,235 m3"),
    ],
    "2.3": [
        ("Traka 400 mm (gore) / 500 mm (dolje) × 2600 mm, PUNE dubine 900 mm = 1,053 m³ "
         "po traci. 4 nosača × 2 trake × 1,053 m³",
         "Traka 500 × 2600 mm, jedinstvene širine po cijeloj dubini (bez proširenja u dnu), "
         "PUNE dubine 900 mm = 1,170 m³ po traci. 4 nosača × 2 trake × 1,170 m³"),
        ("Traka se betonira punom dubinom rova — vlastita težina trake je dio dokaza "
         "sigurnosti na podizanje i prevrtanje.",
         "Traka se betonira punom dubinom rova i do zida rova — vlastita težina trake je "
         "dio dokaza sigurnosti na podizanje i prevrtanje. Temelj širi u dnu nego u vrhu "
         "tražio bi potkopavanje i u stijeni nije izvodiv, pa je presjek jedinstven."),
    ],
    "2.4": [
        ("Obrada vidljivih strana i gornje površine temeljnih traka cementnim malterom u "
         "nagibu 1,0 %, zaglađeno do crnog sjaja.",
         "Obrada vidljivih strana i gornje površine temeljnih traka cementnim malterom u "
         "nagibu 1,0 %, zaglađeno do crnog sjaja. 0,50 × 2,60 × 8 traka = 10,4 m²."),
    ],
}

# (jedinica, stara količina, nova količina)
QTY = {"2.2a": ("m3", 8.94, 9.88), "2.3": ("m3", 8.42, 9.36), "2.4": ("m2", 12, 10.4)}

DROP = "2.2"                        # Zatrpavanje — nema ga (R4)
RENUMBER = [("2.2a", "2.2"), ("2.2b", "2.3"), ("2.3", "2.4"), ("2.4", "2.5"), ("2.5", "2.6")]

COVER = ["UZ PONUDU", "tehničko rješenje", "BAS EN 1993-1-1", "uslov je za počinjanje radova",
         "1,170 m³", "jednom širinom", "10,4 m²"]
GONE = ["≥13,3 kN", "≥16,1 kN", "karakteristična sila čupanja ≥30 kN", "γG,stb = 0,9",
         "qp ≥ 1,20 kN/m² pri nagibu 45°", "uslov za primopredaju", "Zatrpavanje preostalog",
         "1,053 m³"]


def main():
    if not os.path.exists(XLSX):
        raise SystemExit(f"nema fajla: {XLSX}")
    rows_before = dc.xml_rows(XLSX)
    src = openpyxl.load_workbook(XLSX)
    wb = openpyxl.load_workbook(XLSX)
    before = dc.props(wb)
    ws = wb["LOT 1"]
    log = []

    items = jb.structure(ws)["items"]
    if "2.6" in items and "2.2b" not in items:
        # sekcija 2 je već prenumerisana (2.1…2.6) — izmjene su u fajlu
        print(f"ništa za uraditi — sve izmjene su već u fajlu:\n  {XLSX}")
        return
    for n, pairs in EDITS.items():
        c = ws.cell(items[n], 2)
        c.value, done = jb.fix_text(c.value, pairs, f"LOT 1 {n}", idempotent=True)
        if done:
            log.append(f"LOT 1 {n}: {done} zamjena")

    for n, (unit, old, new) in QTY.items():
        r = items[n]
        if ws.cell(r, 3).value != unit:
            raise SystemExit(f"LOT 1 {n}: jedinica {ws.cell(r, 3).value!r} != {unit!r}")
        v = ws.cell(r, 4).value
        if v == new:
            continue
        if v != old:
            raise SystemExit(f"LOT 1 {n}: količina {v}, očekivano {old} ili {new}")
        ws.cell(r, 4).value = new
        log.append(f"LOT 1 {n}: količina {old} -> {new} {unit}")

    dropped = 0
    if DROP in items:
        dropped = r = items[DROP]
        jb.delete_row(ws, r)
        log.append(f"LOT 1 {DROP}: red {r} obrisan (Zatrpavanje — nema ga)")
        items = jb.structure(ws)["items"]
        for old, new in RENUMBER:
            ws.cell(items[old], 1).value = new
        log.append("LOT 1: sekcija 2 prenumerisana " +
                   ", ".join(f"{o}->{n}" for o, n in RENUMBER))

    if not log:
        print(f"ništa za uraditi — sve izmjene su već u fajlu:\n  {XLSX}")
        return

    # brisanje reda pomjera sve ispod njega za jedan, pa se automatske visine iz
    # izvornika prenose kroz to pomjeranje (dc.restore_auto_heights to radi samo za LOT 2)
    for w in wb.worksheets:
        auto = {r for r, (_, custom) in rows_before.get(w.title, {}).items() if not custom}
        if w.title == "LOT 1" and dropped:
            auto = {r if r < dropped else r - 1 for r in auto}
        for r in auto:
            if r in w.row_dimensions:
                w.row_dimensions[r].height = None

    has = "\n".join(str(w.cell(r, 2).value) for w in wb.worksheets
                    for r in range(1, w.max_row + 1) if w.cell(r, 2).value)
    missing = [t for t in COVER if t not in has]
    if missing:
        raise SystemExit(f"nedostaje u predmjeru: {missing}")
    left = [t for t in GONE if t in has]
    if left:
        raise SystemExit(f"trebalo je da nestane iz predmjera: {left}")

    bak = os.path.join(tempfile.gettempdir(), "PRILOG II Sjednica - prije fix_boq_recenzija_a.xlsx")
    shutil.copy2(XLSX, bak)
    wb.save(XLSX)
    after = dc.props(openpyxl.load_workbook(XLSX))
    changed = [k for k in before if before[k] != after.get(k)]
    if changed:
        shutil.copy2(bak, XLSX)
        raise SystemExit(f"snimanje je promijenilo {changed} — izvornik vraćen iz {bak}")
    print(f"snimljeno: {XLSX}\nbackup:    {bak}")
    for line in log:
        print("  " + line)


if __name__ == "__main__":
    main()
