# -*- coding: utf-8 -*-
"""2026-09-11 — prava orijentacija, FN polje prema JZ, nosač 4x3L (odluka Naručioca).

Google Maps: kompleks Sjednica je zakrenut 45° — vrata JI, hladnjak SZ, FN polje JZ,
vanjski ormari SI — pa se zidovi i strane imenuju po stranama svijeta. FN polje: 4
odvojena nosača po 3 modula (1 × 3, položeno), nagib 45°, azimut 225°, u jednom nizu uz
jugozapadnu ogradu (manja „jedra" umjesto jednog polja). Zajednički predmjer
(rfp-hybrid-solar-td, joint_boq.py) kopira ovaj fajl, pa izmjene važe i za Hamziće;
tamo HAMZICI_*_EDITS mijenjaju samo činjenice lokacije.

  O1  LOT 1 · 1.1   4 nosača 1 × 3 (položeno), azimut 225°, gabariti, sile vjetra, stringovi
  O2  LOT 1 · 1.2   sile po nosaču i traci; alternativno sidrenje uz trake iz 2.3
  O3  LOT 1 · 1.4   uzemljenje sva četiri nosača
  O4  LOT 1 · 2.1–2.3   8 traka 400/500 × 900 × 2600 mm (pvsim/stands.py, 4x3L)
  O5  LOT 1 · količine: 1.1/1.2/1.4 = 4 kpl; 2.1 9,88; 2.2 0,94; 2.2a 8,94; 2.2b 0,52;
      2.3 8,42 m³ (cad/design.json foundation)
  O6  LOT 2 · 4.1, 4.2, 4.5–4.8, 4.10   zidovi i strane po stranama svijeta

Idempotentno, obrazac kao fix_boq_dc_aux.py (bez umetanja redova): drugi prolaz ništa ne
mijenja i ne snima fajl; zaglavlje, postavke štampe, svojstva dokumenta i visine redova
provjeravaju se poslije snimanja, a rezervna kopija ide u privremeni folder.

    python fix_boq_sw_4x3l.py
"""
import os
import shutil
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fix_boq_dc_aux as dc  # noqa: E402  (XLSX, jb, xml_rows, props, assert_no_loss)

XLSX, jb = dc.XLSX, dc.jb

EDITS = {
    "LOT 1": {
        "1.1": [
            ("— 3 nosača, svaki za 4 fotonaponska modula (2 reda × 2 kolone, portret); geometrija "
             "margina/prepusta preuzeta iz Huawei Standard A-Shaped Support 3.0 konvencije (PVM "
             "Tabela 4-20), ali poprečna greda izrađena po mjeri za 2 kolone modula, ne za "
             "kataloških 6",
             "— 4 odvojena nosača u jednom nizu, svaki za 3 fotonaponska modula (3 reda × 1 "
             "kolona, položeno); odvojeni nosači umjesto jednog velikog polja smanjuju silu vjetra "
             "po konstrukciji (zahtjev Naručioca); geometrija margina/prepusta preuzeta iz Huawei "
             "Standard A-Shaped Support 3.0 konvencije (PVM Tabela 4-20), ali poprečna greda "
             "izrađena po mjeri za 1 kolonu položenih modula, ne za kataloških 6"),
            (" - inklinacija: fiksno 45°, orijentacija JUG (azimut 180°).",
             " - inklinacija: fiksno 45°, orijentacija JUGOZAPAD (azimut 225°), niz paralelan sa "
             "jugozapadnom ogradom (Naručilac 11.09.2026)."),
            (" - gabariti PV polja po nosaču (2 reda × 2 modula, portret): širina polja 2305 mm "
             "(poprečna greda 2918 mm, bočni prepust 306,5 mm), dužina polja po nagibu 4576 mm "
             "(2 × 2278 mm, nepromijenjeno), horizontalna projekcija 3236 mm pri nagibu 45°; donja "
             "ivica panela na +0,50 m, gornja ivica na +3,74 m od nivoa terena",
             " - gabariti PV polja po nosaču (3 reda × 1 modul, položeno): širina polja 2278 mm "
             "(poprečna greda 2891 mm, bočni prepust 306,5 mm), dužina polja po nagibu 3442 mm "
             "(3 × 1134 mm + 2 × 20 mm), horizontalna projekcija 2434 mm pri nagibu 45°; donja "
             "ivica panela na +0,50 m, gornja ivica na +2,93 m od nivoa terena; razmak između "
             "nosača 400 mm, dužina niza 10 312 mm"),
            ("smještaj: nosači se temelje IZVAN ograđenog platoa, južno od ograde, u pojasu širine "
             "cca 1950 mm;",
             "smještaj: nosači se temelje IZVAN ograđenog platoa, jugozapadno od ograde (JZ "
             "strana), cca 0,40 m od ograde;"),
            ("gornja (sjeverna) ivica panela je 1,64 m iznad kote ograde h=2,10 m.",
             "gornja (sjeveroistočna) ivica panela je 0,83 m iznad kote ograde h=2,10 m."),
            ("Površina izloženosti vjetru 10,55 m² po nosaču (2305 × 4576 mm). Projektne sile po "
             "nosaču (GSN, γQ=1,5 / γG,fav=0,9): podizanje ≥18,1 kN, horizontalna sila ≥13,4 kN, "
             "moment prevrtanja ≥42,6 kNm.",
             "Površina izloženosti vjetru 7,84 m² po nosaču (2278 × 3442 mm). Projektne sile po "
             "nosaču (GSN, γQ=1,5 / γG,fav=0,9): podizanje ≥13,3 kN, horizontalna sila ≥10,0 kN, "
             "moment prevrtanja ≥25,7 kNm."),
            ("(string 1 — nosači PV-1 i PV-2; string 2 — nosači PV-2 i PV-3)",
             "(string 1 — nosači PV-1 i PV-2; string 2 — nosači PV-3 i PV-4)"),
        ],
        "1.2": [
            ("Ukupna projektna sila podizanja po nosaču ≥18,1 kN (GSN); sila po traci od momenta "
             "prevrtanja ≥26,6 kN pri razmaku traka 1600 mm",
             "Ukupna projektna sila podizanja po nosaču ≥13,3 kN (GSN); sila po traci od momenta "
             "prevrtanja ≥16,1 kN pri razmaku traka 1600 mm"),
            ("dozvoljena su SAMO uz gravitacioni temelj zapremine ≥0,75 m³ po nosaču (tendovano "
             "0,81 m³ po nosaču, v. Tačku 2.3)",
             "dozvoljena su SAMO uz gravitacioni temelj iz Tačke 2.3 (dvije trake po nosaču, "
             "najmanje 0,74 m³ po traci; tendovano 1,053 m³ po traci)"),
        ],
        "1.4": [("povezivanje sva tri nosača", "povezivanje sva četiri nosača")],
        "2.1": [("Rov širine 550 mm, dubine 950 mm, dužine 3300 mm. 3 nosača × 2 trake × 1,725 m3",
                 "Rov širine 500 mm, dubine 950 mm, dužine 2600 mm. 4 nosača × 2 trake × 1,235 m3")],
        "2.2": [("3 nosača × 2 trake × 0,149 m3", "4 nosača × 2 trake × 0,117 m3")],
        "2.2b": [("3 nosača × 2 trake × 0,091 m3", "4 nosača × 2 trake × 0,065 m3")],
        "2.3": [("Traka 450 mm (gore) / 550 mm (dolje) × 3300 mm, PUNE dubine 900 mm = 1,485 m³ "
                 "po traci. 3 nosača × 2 trake × 1,485 m³",
                 "Traka 400 mm (gore) / 500 mm (dolje) × 2600 mm, PUNE dubine 900 mm = 1,053 m³ "
                 "po traci. 4 nosača × 2 trake × 1,053 m³")],
    },
    "LOT 2": {
        "4.1": [
            ("unos kroz kapiju na sredini istočne strane ograde",
             "unos kroz kapiju na sredini jugoistočne strane ograde"),
            ("sa JUŽNE i 720 mm sa SJEVERNE strane", "sa JUGOZAPADNE i 720 mm sa SJEVEROISTOČNE strane"),
            ("1155 mm sa ISTOČNE strane. Sa ZAPADNE strane je hladnjak",
             "1155 mm sa JUGOISTOČNE strane. Sa SJEVEROZAPADNE strane je hladnjak"),
        ],
        "4.2": [("smještaj u JUGOISTOČNI ugao kontejnera, prema crtežu M-01",
                 "smještaj u JUŽNI ugao kontejnera (uz JZ i JI zid), prema crtežu M-01")],
        "4.5": [("do žaluzine 600 × 600 mm kroz ZAPADNI zid;",
                 "do žaluzine 600 × 600 mm kroz SJEVEROZAPADNI (SZ) zid;"),
                ("kroz ZAPADNI zid kontejnera do izlazne žaluzine",
                 "kroz SJEVEROZAPADNI (SZ) zid kontejnera do izlazne žaluzine")],
        "4.6": [("žaluzina se ugrađuje u ZAPADNI zid kontejnera, na osi radijatora agregata. Topli "
                 "zrak i izduv se NE smiju izbacivati prema SJEVERNOJ strani,",
                 "žaluzina se ugrađuje u SJEVEROZAPADNI (SZ) zid kontejnera, na osi radijatora "
                 "agregata. Topli zrak i izduv se NE smiju izbacivati prema SJEVEROISTOČNOJ (SI) "
                 "strani,")],
        "4.7": [("žaluzina se ugrađuje u SJEVERNI zid kontejnera, na istočnom kraju,",
                 "žaluzina se ugrađuje u SJEVEROISTOČNI (SI) zid kontejnera, na jugoistočnom kraju,"),
                ("Sjeverna strana je zasjenjena i daje najhladniji usisni zrak",
                 "Sjeveroistočna strana je zasjenjena veći dio dana i daje najhladniji usisni zrak"),
                ("Žaluzina se postavlja istočno od vanjskih ormara",
                 "Žaluzina se postavlja jugoistočno od vanjskih ormara")],
        "4.8": [("ventilator se ugrađuje u ISTOČNI zid kontejnera, u gornjoj zoni (donja ivica cca "
                 "1,75 m), sjeverno od ulaznih vrata.",
                 "ventilator se ugrađuje u JUGOISTOČNI (JI) zid kontejnera, u gornjoj zoni (donja "
                 "ivica cca 1,75 m), sjeveroistočno od ulaznih vrata.")],
        "4.10": [("uspon uz ZAPADNI zid kontejnera", "uspon uz SJEVEROZAPADNI (SZ) zid kontejnera")],
    },
}

# (jedinica, stara količina, nova količina) — cad/design.json foundation, 8 traka
QTY = {"LOT 1": {"1.1": ("kpl", 3, 4), "1.2": ("kpl", 3, 4), "1.4": ("kpl", 3, 4),
                 "2.1": ("m3", 10.35, 9.88), "2.2": ("m3", 0.89, 0.94),
                 "2.2a": ("m3", 9.45, 8.94), "2.2b": ("m3", 0.54, 0.52),
                 "2.3": ("m3", 8.91, 8.42)}}

COVER = ["4 odvojena nosača", "azimut 225°", "2434 mm", "+2,93 m", "25,7 kNm", "16,1 kN",
         "1,053 m³", "PV-3 i PV-4", "sva četiri nosača", "SJEVEROZAPADNI (SZ) zid",
         "SJEVEROISTOČNI (SI) zid", "JUGOISTOČNI (JI) zid", "JUŽNI ugao kontejnera"]


def main():
    if not os.path.exists(XLSX):
        raise SystemExit(f"nema fajla: {XLSX}")
    rows_before = dc.xml_rows(XLSX)
    src = openpyxl.load_workbook(XLSX)
    wb = openpyxl.load_workbook(XLSX)
    before = dc.props(wb)
    log = []
    for sheet, edits in EDITS.items():
        ws = wb[sheet]
        items = jb.structure(ws)["items"]
        for n, pairs in edits.items():
            c = ws.cell(items[n], 2)
            c.value, done = jb.fix_text(c.value, pairs, f"{sheet} {n}", idempotent=True)
            if done:
                log.append(f"{sheet} {n}: {done} zamjena")
    for sheet, q in QTY.items():
        ws = wb[sheet]
        items = jb.structure(ws)["items"]
        for n, (unit, old, new) in q.items():
            r = items[n]
            if ws.cell(r, 3).value != unit:
                raise SystemExit(f"{sheet} {n}: jedinica {ws.cell(r, 3).value!r} != {unit!r}")
            v = ws.cell(r, 4).value
            if v == new:
                continue
            if v != old:
                raise SystemExit(f"{sheet} {n}: količina {v}, očekivano {old} ili {new}")
            ws.cell(r, 4).value = new
            log.append(f"{sheet} {n}: količina {old} -> {new} {unit}")
    if not log:
        print(f"ništa za uraditi — sve izmjene su već u fajlu:\n  {XLSX}")
        return

    dc.restore_auto_heights(wb, rows_before, None)
    dc.assert_no_loss(wb, src)
    has = "\n".join(str(w.cell(r, 2).value) for w in wb.worksheets
                    for r in range(1, w.max_row + 1) if w.cell(r, 2).value)
    missing = [t for t in COVER if t not in has]
    if missing:
        raise SystemExit(f"nedostaje u predmjeru: {missing}")

    bak = os.path.join(tempfile.gettempdir(), "PRILOG II Sjednica - prije fix_boq_sw_4x3l.xlsx")
    shutil.copy2(XLSX, bak)
    wb.save(XLSX)
    after = dc.props(openpyxl.load_workbook(XLSX))
    changed = [k for k in before if before[k] != after.get(k)]
    if changed:
        shutil.copy2(bak, XLSX)
        raise SystemExit(f"snimanje je promijenilo {changed} — izvornik vraćen iz {bak}")
    rows_after, bad = dc.xml_rows(XLSX), []
    for name, flags in rows_before.items():
        for r, (ht, custom) in flags.items():
            got = rows_after.get(name, {}).get(r, (None, False))
            if custom and (not got[1] or float(got[0]) != float(ht)):
                bad.append(f"{name} red {r}: fiksna visina {ht} -> {got}")
            if not custom and got[1]:
                bad.append(f"{name} red {r}: automatska visina postala fiksna {got}")
    if bad:
        shutil.copy2(bak, XLSX)
        raise SystemExit("visine redova:\n  " + "\n  ".join(bad[:10]))
    print(f"snimljeno: {XLSX}\nbackup:    {bak}")
    for line in log:
        print("  " + line)


if __name__ == "__main__":
    main()
