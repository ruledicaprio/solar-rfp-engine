# -*- coding: utf-8 -*-
"""Joint bill of quantities (PRILOG II) for BS Sjednica (Bileća) and BS Hamzići (Čitluk).

The single-site Sjednica workbook (bht-sjednica-final-review/TD-OUTPUT, read-only)
is the template. Every LOT sheet here is a copy_worksheet() of a source sheet, so
styles, merges, column widths, row heights and print setup come from the source;
only the changes below are made.

    LOT 1 Sjednica, LOT 1 Hamzići   PV stand structures - one design, the same items
                                    and quantities at both sites; the site facts in
                                    1.1 and 2.3 differ (HAMZICI_LOT1_EDITS) and Hamzići
                                    adds 1.6, the existing earth rings under the strips
    LOT 2 Sjednica, LOT 2 Hamzići   genset in the existing container - Hamzići differs
                                    only where the site differs (HAMZICI_*, NEW_ITEMS)
    REKAPITULACIJA LOT n            the LOT's two site totals, its own discount and
                                    17 % VAT, notes, date and signature

One workbook per LOT (Investor, 12.09.2026): the award is per LOT and a bidder may
offer only one, so each LOT is a complete price form of its own - its two site
sheets and its REKAPITULACIJA. The source REKAPITULACIJA / NAPOMENA / Datum blocks
below each LOT subtotal are dropped; the recap finds every subtotal cell by its
label, never by a hard-coded row.

Source defects fixed on the way (listed again in the run summary):
  * fit-to-width printing is switched on - the source carries fitToWidth=1 but
    fitToPage is off, so it printed at 100 % (LOT 2 across two page columns);
  * rows with wrapped text are left to auto-fit - the stored heights are stale
    (item 1.1 prints two of its ~40 lines in a fixed 24 pt row);
  * LOT 2 columns E/F take LOT 1's widths - 45.81 / 5.0 put the totals on a
    second page and show them as ###;
  * the "UKUPNO 1" label moves from column A (5.7 wide, printed as "NELE:") to B;
  * item-row cells without border / number format get the column's item style;
  * LOT 2 item 4.5 pointed to the outlet louvre as "Tačka 4.7" (the intake) - it is
    4.6, fixed on both site sheets (BOTH_SITES_EDITS).

    python joint_boq.py          writes paths.PRILOG2_LOT, one file per LOT (TD_OUT redirects)
"""
import os
import re
import sys
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import paths  # noqa: E402

SRC = os.path.join(paths.SITES["sjednica"]["folder"], "TD-OUTPUT",
                   "3.1 PRILOG II TD - predmjer Sjednica Bileca.xlsx")
OUT = paths.PRILOG2_LOT                  # {lot: path} - one workbook per LOT

SITE = {
    "sjednica": {"up": "BS SJEDNICA", "title": "BS SJEDNICA (BILEĆA)", "label": "BS Sjednica (Bileća)"},
    "hamzici": {"up": "BS HAMZIĆI", "title": "BS HAMZIĆI (ČITLUK)", "label": "BS Hamzići (Čitluk)"},
}
LOTS = ("LOT 1", "LOT 2")
SHEET = {(lot, s): f"{lot} {paths.SITES[s]['name']}" for lot in LOTS for s in SITE}
SITE_SHEETS = {lot: [SHEET[lot, s] for s in SITE] for lot in LOTS}
LOT_SHEETS = [n for lot in LOTS for n in SITE_SHEETS[lot]]
REKAP = {lot: f"REKAPITULACIJA {lot}" for lot in LOTS}
SHEET_ORDER = {lot: SITE_SHEETS[lot] + [REKAP[lot]] for lot in LOTS}

SRC_SITE = " SJEDNICA, BILEĆA"           # how the source title rows name the site
JOINT = "\nBS SJEDNICA I BS HAMZIĆI"     # ... and the joint tender (own line: fits any sheet)
PRINT_TITLES = "8:9"                     # column numbers + column headings
NUM = "#,##0.00"

ITEM_RE = re.compile(r"^\d+\.\d+[a-z]?$")   # "4.1", "4.10", "2.2a" - matched exactly
SUB_RE = re.compile(r"^UKUPNO (\d+)\b")      # section subtotal "UKUPNO 5 — ..."
LOT_RE = re.compile(r"^UKUPNO LOT (\d) —")   # LOT subtotal "UKUPNO LOT 2 — ..." (not the
#                                              source recap's "UKUPNO LOT 1 (bez PDV-a)")
LINE = '=IF(AND(D{r}<>"",E{r}<>""),D{r}*E{r},"")'   # source line-total pattern

# ---- REKAPITULACIJA labels, one recap per LOT (the check script imports these) --
LBL_SITE = {(lot, s): f"{SITE[s]['label']} — {lot} bez PDV-a [KM]:" for lot in LOTS for s in SITE}
LBL_LOT_TOTAL = {lot: f"UKUPNO {lot} (BS Sjednica + BS Hamzići) bez PDV-a [KM]:" for lot in LOTS}
LBL_POPUST = "Popust [%]:"
LBL_DISC = {lot: f"UKUPNO {lot} sa popustom bez PDV-a [KM]:" for lot in LOTS}
LBL_VAT = "Iznos PDV-a (17 %) [KM]:"
LBL_GRAND = {lot: f"UKUPNO {lot} sa popustom i PDV-om [KM]:" for lot in LOTS}

# ---- BS Hamzići, LOT 2: what differs from Sjednica -------------------------
# Text edits are exact (old, new) replacements inside the Sjednica item text; each
# old fragment must occur exactly once, so everything else in the item is kept.
HAMZICI_EDITS = {
    "4.5": [
        # Sjednica wording since fix_boq_sw_4x3l.py (walls by true direction, 11.09.2026)
        ("Kanal je PRELAZNI KOMAD od prirubnice hladnjaka do žaluzine 600 × 600 mm kroz "
         "SJEVEROZAPADNI (SZ) zid; hladnjak je 60 mm od zida, pa je razvijena površina ≈1,0 m².",
         "Kanal je PLENUM od prirubnice hladnjaka do POSTOJEĆIH OTVORA demontiranog klima-uređaja "
         "Stulz (Tačka 5.15) u sredini JUGOISTOČNOG (JI) zida kontejnera, sa prirubnicama, "
         "fleksibilnim spojem prema hladnjaku (Tačka 4.4) i brtvljenjem; razvijena površina ≈1,5 m²."),
        ("RASPORED (OBAVEZNO, prema crtežu M-01): kanal se vodi najkraćim putem od radijatora kroz "
         "SJEVEROZAPADNI (SZ) zid kontejnera do izlazne žaluzine iz Tačke 4.6.",   # after T1
         "RASPORED (OBAVEZNO, prema crtežu H-04): kanal se vodi najkraćim putem od radijatora do "
         "postojećih otvora klima-uređaja Stulz u sredini JUGOISTOČNOG (JI) zida kontejnera, u "
         "kojima se ugrađuje izlazna žaluzina iz Tačke 4.6."),
    ],
    "4.6": [
        ("Isporuka i montaža fiksne žaluzine na kraju kanala za odvod toplog zraka, dimenzija "
         "600 × 600 mm, sa zaštitnom mrežicom, komplet sa montažnim materijalom.",
         "Prilagođenje postojećih otvora demontiranog klima-uređaja Stulz (Tačka 5.15) u sredini "
         "JUGOISTOČNOG (JI) zida kontejnera za izlaz toplog zraka sa hladnjaka agregata: spajanje i/ili "
         "proširenje otvora na žaluzinu bruto površine ≥0,36 m² (npr. 600 × 600 mm), sa ojačanim "
         "okvirom u zidnom panelu; isporuka i montaža fiksne žaluzine na kraju kanala za odvod "
         "toplog zraka, sa zaštitnom mrežicom protiv insekata i ptica, komplet sa montažnim "
         "materijalom; vani hauba 700 × 600 mm od pocinčanog lima, zatvorenih bočnih strana i dna, "
         "sa rešetkom na vrhu (≈+1,30 m) — topli zrak se usmjerava NAVIŠE, ne prema FN polju; hauba "
         "najmanje 1,0 m od nogu stuba; neiskorišteni dio postojećih otvora zatvoriti sendvič "
         "panelom d=60 mm istovjetnim zidu kontejnera, sa brtvljenjem."),
        ("RASPORED (OBAVEZNO, prema crtežu M-01): žaluzina se ugrađuje u SJEVEROZAPADNI (SZ) zid "
         "kontejnera, na osi radijatora agregata. Topli zrak i izduv se NE smiju izbacivati prema "
         "SJEVEROISTOČNOJ (SI) strani, gdje se nalaze postojeći vanjski ormari ICC360-HA1-C1 i "
         "MTS9302A i gdje se nalazi usis svježeg zraka.",
         "RASPORED (OBAVEZNO, prema crtežu H-04): žaluzina se ugrađuje u JUGOISTOČNI (JI) zid "
         "kontejnera, na mjestu postojećih otvora klima-uređaja Stulz u sredini zida, sa haubom "
         "koja topli zrak usmjerava naviše. Topli zrak i izduv se NE smiju izbacivati prema "
         "SJEVEROZAPADNOJ (SZ) strani, gdje su ulazna vrata kontejnera, prema JUGOZAPADNOJ (JZ) "
         "strani, gdje su vanjski ormari ICC360-HA1-C1 i MTS i FN polje, niti prema usisu svježeg "
         "zraka na SJEVEROISTOČNOM (SI) zidu (Tačka 4.7)."),
    ],
    "4.7": [
        ("RASPORED (OBAVEZNO, prema crtežu M-01): žaluzina se ugrađuje u SJEVEROISTOČNI (SI) zid "
         "kontejnera, na jugoistočnom kraju, sa donjom ivicom na cca 0,30 m od poda. Sjeveroistočna "
         "strana je zasjenjena veći dio dana i daje najhladniji usisni zrak, čime se poboljšava "
         "hlađenje agregata. Žaluzina se postavlja jugoistočno od vanjskih ormara "
         "ICC360-HA1-C1/MTS9302A kako se ne bi usisavao topli zrak sa njih.",
         "RASPORED (OBAVEZNO, prema crtežu H-04): žaluzina se ugrađuje u SJEVEROISTOČNI (SI) zid "
         "kontejnera, uz alternator (sjeverozapadni kraj agregata), jugoistočno od korita "
         "spremnika, sa donjom ivicom na cca 0,30 m od poda; sjeveroistočna strana je zasjenjena "
         "veći dio dana."),
        ("spremnika ≥3 m.", "spremnika ≥3 m (H-04: ≥3,2 m do završetka izduva i do oduška "
         "spremnika)."),
    ],
    "4.10": [
        ("izvedena izvan kontejnera i završena IZNAD KROVA, usmjereno naviše, sa kapom protiv upada "
         "padavina.",
         "izvedena izvan kontejnera HORIZONTALNO kroz JUGOISTOČNI (JI) zid kontejnera, iz "
         "jugozapadnog (JZ) prolaza, zapadno od haube, na visini ≈+2,30 m, usmjereno prema "
         "jugoistoku, sa kapom protiv upada padavina. Prodor kroz zid izvesti sa termički "
         "izolovanom zaštitnom čahurom."),
        ("od prigušivača do izlaza iznad krova, do 4 m.",
         "od prigušivača do izlaza kroz JI zid, do 4 m."),
        ("TRASA: fleksibilni spoj i prigušivač neposredno iza motora, uspon uz SJEVEROZAPADNI (SZ) "
         "zid kontejnera, završetak IZNAD KROVA usmjeren naviše, sa hvatačem iskri.",
         "TRASA: fleksibilni spoj neposredno iza motora, prigušivač u jugozapadnom (JZ) prolazu, "
         "horizontalni izlaz kroz JUGOISTOČNI (JI) zid kontejnera na ≈+2,30 m, zapadno od haube, "
         "sa završetkom ≥0,40 m od zida — dalje od FN polja, ormara i ulaznih vrata, sa hvatačem "
         "iskri. Završetak iznad krova NIJE moguć: platforma antenskog stuba na +3,0 m nalazi se "
         "iznad krova kontejnera."),
        ("protutlak ≈1,9 kPa", "protutlak ≈1,6 kPa"),          # hamzici review/07 C.4
    ],
    # Site facts (hamzici-hybrid-solar/cad/*.json, review/07-proracuni.md, joint Prilog I).
    "3.1": [
        ("Snaga prema ISO 3046 za nadmorsku visinu preko 1000 m (lokacija 1076 m n.v.);",
         "Snaga prema ISO 3046 za nadmorsku visinu lokacije 493 m n.v. i temperaturu okoline "
         "+40 °C: faktor 0,931, standby 13,4 kW, prime 12,3 kW;"),
        ("(≈82 % derativane prime snage na 1076 m n.v. i +40 °C, koja iznosi 11,6 kW)",
         "(≈77 % derativane prime snage na 493 m n.v. i +40 °C, koja iznosi 12,3 kW)"),
    ],
    "4.1": [
        ("unos kroz kapiju na sredini jugoistočne strane ograde (svijetla širina cca 1,00 m) i "
         "ulazna vrata kontejnera 900 × 2000 mm;",
         "unos kroz kapiju na SJEVEROZAPADNOJ (SZ) strani ograde (svijetla širina 1,30 m) i ulazna "
         "vrata kontejnera na SZ zidu, 1,00 × 2,15 m — skid 620 mm kroz vrata svijetle širine "
         "990 mm, pravo po osi; najprije agregat, zatim spremnik;"),
        ("SERVISNI PROSTOR ( prema crtežu M-01): agregat se postavlja CENTRIRANO u slobodnom prostoru "
         "kontejnera, sa najmanje 720 mm sa JUGOZAPADNE i 720 mm sa SJEVEROISTOČNE strane (520 mm na "
         "dijelu gdje je GRO) te 1155 mm sa JUGOISTOČNE strane. Sa SJEVEROZAPADNE strane je hladnjak, "
         "koji izduvava u kanal kroz zid i ne servisira se s te strane",
         "SERVISNI PROSTOR (prema crtežu H-04): os agregata SZ–JI u sredini kontejnera; slobodan "
         "prostor: jugozapadna (JZ) strana 0,78 m, sjeveroistočna (SI) strana 0,78 m, kraj prema SZ "
         "0,97 m do čela GRO; od vrata do korita spremnika 0,95 m slobodno. Sa JUGOISTOČNE (JI) "
         "strane je hladnjak, koji preko plenuma od 0,12 m izduvava kroz postojeće otvore "
         "klima-uređaja Stulz u sredini JI zida, u haubu koja topli zrak usmjerava naviše (Tačke 4.5 "
         "i 4.6) i ne servisira se s te strane. Ponuđač potvrđuje da su servisna mjesta agregata "
         "dostupna sa JZ i SI strane i sa SZ kraja"),
    ],
    "4.8": [
        # the source (fix_boq_dc_aux.py D2) makes it a 48 V DC EC fan off the DC razvod that
        # the genset controller switches off while the genset runs; Hamzići adds: extract fan
        ("aksijalnog EC ventilatora 48 V DC za prinudnu ventilaciju prostora agregata",
         "aksijalnog EC ventilatora 48 V DC, izvlačnog — izbacuje zrak iz prostora, za prinudnu "
         "ventilaciju prostora agregata"),
        ("RASPORED (prema crtežu M-01)", "RASPORED (prema crtežu H-04)"),
        ("ventilator se ugrađuje u JUGOISTOČNI (JI) zid kontejnera, u gornjoj zoni (donja ivica "
         "cca 1,75 m), sjeveroistočno od ulaznih vrata.",
         "ventilator se ugrađuje u JUGOISTOČNI (JI) zid kontejnera, u gornjoj zoni, sjeveroistočno "
         "od haube (Tačka 4.6)."),
    ],
    "4.2": [
        ("smještaj u JUŽNI ugao kontejnera (uz JZ i JI zid), prema crtežu M-01",
         "smještaj u SJEVERNI ugao kontejnera (uz SZ i SI zid), prema crtežu H-04 — prihvatno "
         "korito 1150 × 640 mm uz SZ i SI zid; izduvni cjevovod je na jugozapadnoj strani agregata"),
        ("odušna cijev izvedena IZVAN kontejnera, otvor zaštićen metalnom mrežicom;",
         "odušna cijev izvedena IZVAN kontejnera kroz SZ zid, sjeveroistočno od vrata, do stojeće "
         "cijevi između kontejnera i SZ ograde (završetak +2,80 m) — principijelno prema H-04, "
         "konačno prema elaboratu zaštite od požara; najmanje 3 m od završetka izduva i od usisa "
         "zraka; otvor zaštićen metalnom mrežicom;"),
    ],
    "4.12": [
        ("Za trasu do 5 m iz Tačke 4.10 i vanjski prečnik izolacije ≈165 mm razvijena površina je "
         "≈3,0 m².",
         "Za trasu iz Tačke 4.10 (≈1 m do prigušivača u JZ prolazu, zatim ≈2 m do JI zida) i "
         "vanjski prečnik izolacije ≈165 mm razvijena površina je ≈2,5 m²."),
    ],
    "5.3": [
        ("na dubini od 1,00 m obzirom da se radi o planinskoj lokaciji (izbjegavanje efekta "
         "zaleđivanja tla),",
         "na dubini od 1,00 m radi mehaničke zaštite i zaštite od smrzavanja tla (jedan opis za obje "
         "lokacije); tlo je krš, stijena blizu površine, pa iskop može biti u stijeni,"),
    ],
    "5.6": [("vanjski sistem zaštite od munje (antenski stub h=38 m)",
             "vanjski sistem zaštite od munje (rešetkasti antenski stub h=32 m, platforme na "
             "+3,0 / +12,0 / +30,0 m)"),
            ("Ormar zidni sa nosačima, orijentacionih dimenzija 0,60 × 0,25 × 0,80 m (Š×D×V)",
             "Ormar zidni sa nosačima, za SZ zid jugozapadno od ulaznih vrata (širina zida "
             "0,595 m), širine ≤0,50 m — orijentacionih dimenzija npr. 0,50 × 0,25 × 0,80 m "
             "(Š×D×V)")],
    "5.11": [("vanjski sistem zaštite od munje i antenski stub h=38 m",
              "vanjski sistem zaštite od munje i rešetkasti antenski stub h=32 m (platforme na "
              "+3,0 / +12,0 / +30,0 m)")],
    "5.13": [("zbog dužine DC trase od 25 m.", "zbog dužine DC trase od ≈20 m u jednom smjeru."),
             ("i crtež E-01)", "i crtež H-05)")],
    # Huawei ICC360-HA1-C1 and the MTS stand outdoors on the slab, JZ side, behind the PV row.
    "5.5": [
        ("Isporuka i polaganje veza na dionici hibridni sistem — DEA:",
         "Isporuka i polaganje veza na dionici hibridni sistem — DEA; ormari Huawei ICC360-HA1-C1 "
         "i MTS stoje na ploči na JZ strani, iza FN polja (principijelno), a veza sa GRO vodi "
         "≈3–5 m kroz JZ zid kontejnera:"),
        ("energetski kabl dužine do 15 m", "energetski kabl dužine ≈3–5 m"),
        ("signalni kabl dužine do 15 m", "signalni kabl dužine ≈3–5 m"),
        ("komunikacioni Ethernet kabl dužine do 15 m", "komunikacioni Ethernet kabl dužine ≈3–5 m"),
    ],
    # The quantity basis written in the item text follows the new quantity.
    "5.1": [("do 15 m1 i krak uzemljivačke trake do nosača, ukupno do 30 m1",
             "do 10 m1 i krak uzemljivačke trake do nosača, ukupno do 20 m1")],
    "5.2": [("3 cijevi × 15,00 m", "3 cijevi × 10,00 m")],
    "5.12": [("dužini 25 m iznosi 0,78", "dužini 20 m iznosi 0,62"),   # hamzici review/07 D.2
             ("2 stringa × 2 × 25,00 m", "2 stringa × 2 × 20,00 m")],
}
HAMZICI_QTY = {"5.1": ("m1", 20), "5.2": ("m", 30), "5.3": ("m", 30), "5.12": ("m", 80)}

# ---- BS Hamzići, LOT 1: same stands and foundations, only the site facts differ
HAMZICI_LOT1_EDITS = {
    "1.1": [
        # no energy figures in Prilog II: the Sjednica December-yield sentence gives way
        ("Nagib je zadržan zbog decembarskog prinosa — pri podnevnoj visini Sunca 23,6° na 42,94° N "
         "nagib 45° ostvaruje 93 % direktnog zračenja u odnosu na 85 % pri 35°, a decembar je "
         "mjerodavni mjesec za dimenzionisanje autonomnog sistema",
         "Listopadno stablo JJI–JI od stuba (≈7–9 m, 15–20 m, izvan zakupa) ostaje; pri polju "
         "okrenutom prema jugozapadu njegov uticaj je mali (Prilog I, Tačka 3.7)"),
        # Sjednica wording since fix_boq_sw_4x3l.py (4x3L, SW, 11.09.2026)
        ("smještaj: nosači se temelje IZVAN ograđenog platoa, jugozapadno od ograde (JZ strana), "
         "cca 0,40 m od ograde;",
         "smještaj (43,288012° N, 17,624794° E, 493 m n.v.): nosači se temelje IZVAN ograđenog "
         "platoa, jugozapadno od ploče i ograde, u pojasu dubine 3300 mm i dužine 12,50 m (JZ "
         "strana), unutar zakupa k.č. 109/1 K.O. Hamzići (12,00 × 12,50 m); trake 350 mm od granice "
         "zakupa i od ploče;"),
        ("gornja (sjeveroistočna) ivica panela je 0,83 m iznad kote ograde h=2,10 m. Ponuđač "
         "provjerava da konstrukcija u cijelosti ostaje unutar zakupljene parcele 16,00 × 9,40 m",
         "gornja (sjeveroistočna) ivica panela je 0,93 m iznad vrha ograde — ograda je h=1,80 m "
         "iznad ploče, odnosno 2,00 m iznad vanjskog terena, koji je uz ploču na −0,20 m (ovjereni "
         "04_Ograda). Ponuđač provjerava da konstrukcija u cijelosti ostaje unutar zakupa "
         "12,00 × 12,50 m"),
        ("prema ovjerenoj projektnoj dokumentaciji lokacije i BAS EN 1991-1-4 sa BiH nacionalnim "
         "aneksom, uz primjenu faktora orografije za izloženi planinski vrh na 1076 m n.v.",
         "prema BAS EN 1991-1-4 sa BiH nacionalnim aneksom — ista vrijednost kao na lokaciji "
         "Sjednica, radi jedne konstrukcije za obje lokacije (iz ovjerenog projekta lokacije, "
         "493 m n.v., izvodi se 0,69–0,96 kN/m²)."),
        ("za konkretnu lokaciju (planinski vrh)", "za konkretnu lokaciju"),
        ("(v. crtež E-01)", "(v. crtež H-05)"),
    ],
    "2.3": [
        ("Klasa XF3 je mjerodavna zbog cikličnog smrzavanja i odmrzavanja u vlažnom stanju na "
         "1076 m n.v.",
         "Klasa XF3: ista specifikacija kao na lokaciji Sjednica, radi jednog opisa za obje "
         "lokacije."),
    ],
}

# Source typos fixed on BOTH site sheets of a LOT, before the Hamzići edits. The Sjednica
# source gets the same fixes from bht-sjednica-final-review/tools/fix_boq_dc_aux.py, so
# they are applied idempotently: a fix already in the source is skipped, not an error.
BOTH_SITES_EDITS = {
    "LOT 2": {
        "4.5": [("do izlazne žaluzine iz Tačke 4.7.",       # 4.7 is the intake louvre
                 "do izlazne žaluzine iz Tačke 4.6.")],
        "4.2": [("ojačanja poda iz Tačke 4.4.",             # the grillage is 4.3, 4.4 the
                 "ojačanja poda iz Tačke 4.3.")],           # radiator's flexible joint
    },
}
BOTH_SITES_NOTE_EDITS = {
    "LOT 2": {"OPŠTE NAPOMENE UZ TAČKU 4:": [
        ("roštilj za raznošenje opterećenja iz Tačke 4.4 OBAVEZAN",
         "roštilj za raznošenje opterećenja iz Tačke 4.3 OBAVEZAN")]},
}

# Rows without an item number, keyed by the start of their column-B text (Hamzići only).
HAMZICI_NOTE_EDITS = {
    "LOT 2": {"OPŠTE NAPOMENE UZ TAČKU 4:": [
        ('ovjerenom projektu lokacije („04 AG dio", tačka 4.4.2.3)',
         "ovjerenom projektu lokacije (AG dio, tačka 4.4.2)")]},
}

# Items replaced as a whole: item -> (expected start of the Sjednica text, unit, qty, new text)
HAMZICI_REPLACE = {
    "5.10": ("Snimanje postojećeg stanja i prevezivanje postojećih strujnih krugova", "kpl", 1,
             "Isporuka, montaža i povezivanje rasvjete i utičnica kontejnera — kontejner je PRAZAN, "
             "bez postojećih električnih instalacija. Obuhvata:\n"
             " - 1 kom LED svjetiljka IP65, svjetlosnog toka ≥1500 lm, 230 V iz GRO; druga svjetiljka "
             "kontejnera je LED svjetiljka 48 V DC sa prekidačem uz vrata iz Tačke 5.16, koja daje "
             "svjetlo i kad agregat ne radi\n"
             " - 2 kom utičnica 230 V / 16 A, IP44\n"
             " - prekidač rasvjete\n"
             " - bezhalogeni kablovi položeni u PVC kanalice, priključak na novi GRO iz Tačke 5.6\n"
             " - postojeća SIGNALNA RASVJETA (rasvjeta prepreke) antenskog stuba NE napaja se iz GRO: "
             "priključuje se na DC razvod −48 V iz Tačke 5.16, svjetiljka prema Tački 5.17 — rasvjeta "
             "prepreke je trajno noćno opterećenje i ne smije biti isključena prilikom rekonfiguracije "
             "napajanja\n"
             " - Ponuđač mjeri stvarnu snagu rasvjete prepreke i dostavlja je Kupcu radi provjere "
             "energetskog bilansa (Prilog I, Tačka 1)\n"
             "Komplet sa spojnim materijalom, oznakama, ispitivanjem i jednopolnom shemom izvedenog "
             "stanja."),
}

# New Hamzići items per LOT: (item, after, unit, qty, text). The row goes right after item
# `after`; its number must be free and fall numerically between `after` and the section's
# next item (minor numbers compared as integers - "5.1" never stands for "5.16"). 5.15 sits
# between 5.14 and the source's 5.16-5.18, which keep their numbers at both sites.
NEW_ITEMS = {
    "LOT 1": [
        ("1.6", "1.5", "kpl", 1,
         "Postojeći uzemljivač na ukrštanjima sa temeljnim trakama: lociranje (detektor, ručni "
         "iskop), otkopavanje i zaštita dva postojeća prstena FeZn 25 × 4 mm na dubini 0,8 m "
         "(kvadrati 7,50 m i 10,00 m oko ploče, 1,05 m i 2,30 m od ivice ploče — ovjereni Plan "
         "uzemljivača 3.6.9); prsten se ne prekida trajno — premješta se ispod ili oko trake ili "
         "premošćuje istim materijalom, spojevi prema EN 62305-3 sa antikorozivnom zaštitom; spoj "
         "sa uzemljenjem nosača; mjerenje otpora rasprostiranja prije i poslije radova, zapisnik."),
    ],
    "LOT 2": [
        ("5.15", "5.14", "kpl", 1,
         "Odspajanje (230 V AC, 48 V DC, signalizacija) i demontaža postojećeg kompaktnog zidnog "
         "klima-uređaja Stulz WDE80 (8 kW) na JUGOISTOČNOM (JI) zidu kontejnera, u sredini, BEZ "
         "otvaranja "
         "rashladnog "
         "kruga; pakovanje, utovar i transport u skladište BH Telecom d.d. na Alipašinom Polju, "
         "Sarajevo, istovar i zapisnik o primopredaji (tip, serijski broj, stanje). Otvori u zidu "
         "ostaju i koriste se za izlaz toplog zraka agregata (Tačka 4.6)."),
    ],
}

EXPECTED_CHANGES = {
    "LOT 1": sorted(set(HAMZICI_LOT1_EDITS) | {n for n, *_ in NEW_ITEMS["LOT 1"]}),
    "LOT 2": sorted(set(HAMZICI_EDITS) | set(HAMZICI_QTY) | set(HAMZICI_REPLACE)
                    | {n for n, *_ in NEW_ITEMS["LOT 2"]}),
}

# ---- NAPOMENA (source LOT 2 notes, adapted to two sites) --------------------
# (exact source note, [(old, new) edits]); None as source = a new note.
NOTES = [
    ("- Nabavka je podijeljena na dva LOT-a. Ponuđač može dostaviti ponudu za JEDAN ili OBA LOT-a. "
     "Ugovor se dodjeljuje po LOT-u, prema kriteriju najniže cijene tehnički zadovoljavajuće ponude.",
     [("Nabavka je podijeljena na dva LOT-a.",
       "Nabavka je podijeljena na dva LOT-a prema vrsti radova: LOT 1 — konstrukcija nosača za "
       "fotonaponske panele, LOT 2 — agregatsko postrojenje (DEA) u postojećem kontejneru. Svaki LOT "
       "obuhvata OBJE lokacije: BS Sjednica (Bileća) i BS Hamzići (Čitluk).")]),
    # new notes: {lot} and {sheets} are filled in per LOT workbook (build_rekap)
    (None,
     "- Ovaj obrazac se odnosi na {lot}. Za drugi LOT popunjava se, potpisuje i ovjerava zaseban "
     "obrazac; ponuđač koji ne pristupa drugom LOT-u ne dostavlja njegov obrazac."),
    (None,
     "- Cijene se upisuju posebno za svaku lokaciju, na listovima {sheets}; iznosi se "
     "automatski prenose u ovu rekapitulaciju."),
    ("- Ugovorne obaveze nastaju po upućivanju pismenog zahtjeva/narudžbe od strane BH Telecom-a.", []),
    ("- Količine u Tačkama 2. i 5. su orijentacione i utvrđene na osnovu podataka iz RFI dokumente. "
     "Konačne količine utvrđuju se elaboratom montaže i geodetskim snimkom, a obračunavaju se po "
     "stvarno izvedenim količinama.",
     [("su orijentacione i utvrđene na osnovu podataka iz RFI dokumente.",
       "su orijentacione, posebno za svaku lokaciju, i utvrđene na osnovu podataka iz RFI "
       "dokumentacije.")]),
    ("- LOT 2 obuhvata montažu i uvezivanje sistema DO POTPUNE GOTOVOSTI.",
     [("DO POTPUNE GOTOVOSTI.", "DO POTPUNE GOTOVOSTI, na obje lokacije.")]),
    ("- Grafički prilozi (situacija postojećeg i budućeg stanja, presjek, dispozicija opreme) dati su "
     "u PRILOGU III tenderske dokumentacije.",
     [("- Grafički prilozi (", "- Grafički prilozi za obje lokacije (")]),
]


# ---- structure helpers (shared with check_boq_recalc.py) --------------------
def line_formula(r):
    return LINE.format(r=r)


def structure(ws):
    """Items, priced items, sections and the LOT subtotal row of a LOT sheet.

    items    {item number: row}, in row order; numbers compare as exact strings
    priced   {item number: True if column F carries the line-total formula}
    sections {section: (first priced row, last priced row, subtotal row)}
    lot_row  row of "UKUPNO LOT n ..."
    """
    items, priced, subs, lot_row = {}, {}, {}, None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and ITEM_RE.match(a.strip()):
            n = a.strip()
            if n in items:
                raise ValueError(f"{ws.title}: item {n} twice (rows {items[n]}, {r})")
            items[n] = r
            f = ws.cell(r, 6).value
            priced[n] = isinstance(f, str) and f.startswith("=IF(AND(")
        for c in (1, 2):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            if LOT_RE.match(v):
                if lot_row:
                    raise ValueError(f"{ws.title}: two LOT subtotals ({lot_row}, {r})")
                lot_row = r
            m = SUB_RE.match(v)
            if m:
                if m.group(1) in subs:
                    raise ValueError(f"{ws.title}: two subtotals for section {m.group(1)}")
                subs[m.group(1)] = r
    sections = {}
    for sec, rsub in subs.items():
        rows = [items[n] for n in items if n.split(".")[0] == sec and priced[n]]
        if not rows:
            raise ValueError(f"{ws.title}: section {sec} has a subtotal but no priced items")
        sections[sec] = (min(rows), max(rows), rsub)
    orphans = [n for n in items if priced[n] and n.split(".")[0] not in sections]
    if orphans or lot_row is None:
        raise ValueError(f"{ws.title}: items without a section subtotal {orphans} / LOT row {lot_row}")
    return {"items": items, "priced": priced, "sections": sections, "lot_row": lot_row}


def find_row(ws, label, col=2):
    hits = [r for r in range(1, ws.max_row + 1) if ws.cell(r, col).value == label]
    if len(hits) != 1:
        raise LookupError(f"{ws.title}: label {label!r} found {len(hits)} times")
    return hits[0]


def replace_once(text, old, new, where):
    n = text.count(old)
    if n != 1:
        raise ValueError(f"{where}: expected 1 occurrence of {old[:60]!r}..., found {n}")
    return text.replace(old, new)


# ---- sheet surgery ----------------------------------------------------------
REF_RE = re.compile(r"(?<![A-Za-z0-9_$!'\"])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![0-9A-Za-z_(!])")


def shift_formula_rows(formula, at, n=1, cross_sheet=False):
    """Excel's own rule for inserting n rows at `at`: same-sheet refs >= at move down.

    A sheet-qualified ref ('LOT 1'!F29) points at another sheet and is left alone; that
    is allowed only with cross_sheet=True, and the caller checks nothing points back in."""
    if "!" in formula and not cross_sheet:
        raise ValueError(f"cross-sheet formula on a LOT sheet: {formula}")

    def rep(m):
        row = int(m.group(4))
        return f"{m.group(1)}{m.group(2)}{m.group(3)}{row + n if row >= at else row}"
    return REF_RE.sub(rep, formula)


def insert_row(ws, at, cross_sheet=False):
    """Insert one row at `at`; row heights, merges and formulas follow the moved rows."""
    dims = {r: ws.row_dimensions[r] for r in list(ws.row_dimensions) if r >= at}
    ws.insert_rows(at)
    for r in dims:
        del ws.row_dimensions[r]
    for r in sorted(dims, reverse=True):
        dims[r].index = r + 1
        ws.row_dimensions[r + 1] = dims[r]
    for mr in ws.merged_cells.ranges:
        if mr.min_row >= at:
            mr.shift(row_shift=1)
        elif mr.max_row >= at:
            raise ValueError(f"{ws.title}: insertion at {at} splits merged range {mr}")
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = shift_formula_rows(c.value, at, cross_sheet=cross_sheet)


def truncate(ws, last):
    """Drop every row below `last` - cells, heights and merges."""
    for key in [k for k in ws._cells if k[0] > last]:
        del ws._cells[key]
    for r in [r for r in ws.row_dimensions if r > last]:
        del ws.row_dimensions[r]
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row > last:
            ws.merged_cells.remove(mr)
        elif mr.max_row > last:
            raise ValueError(f"{ws.title}: merged range {mr} crosses row {last}")


def autofit_rows(ws):
    """Leave rows with wrapped text to auto-fit (no stored height).

    Excel and LibreOffice both size a row without a custom height to its wrapped
    text when the file is opened; merged rows keep their heights (no auto-fit).
    """
    merged = {r for m in ws.merged_cells.ranges for r in range(m.min_row, m.max_row + 1)}
    n = 0
    for r in range(1, ws.max_row + 1):
        if r in merged:
            continue
        if any(isinstance(c.value, str) and c.value.strip() and c.alignment.wrap_text for c in ws[r]):
            if ws.row_dimensions[r].height is not None:
                ws.row_dimensions[r].height = None
                n += 1
    return n


def normalize_item_styles(ws):
    """Item-row cells with no border (or F/E without number format) take the style of the
    same column in the previous complete item row - source formatting slips."""
    s = structure(ws)
    good, fixed = {}, []
    for n, r in s["items"].items():
        if not s["priced"][n]:
            continue
        cols = []
        for col in range(1, 7):
            c = ws.cell(r, col)
            b = c.border
            bordered = any(getattr(b, side).style for side in ("left", "right", "top", "bottom"))
            if bordered and (col < 5 or c.number_format == NUM):
                good[col] = c
            elif col in good:
                c._style = copy(good[col]._style)
                cols.append(get_column_letter(col))
        if cols:
            fixed.append(f"{n}:{''.join(cols)}")
    return fixed


def wrap_title(cell):
    al = copy(cell.alignment)
    al.wrap_text = True
    al.vertical = "center"
    cell.alignment = al


def print_setup(ws, last):
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:F{last}"


# ---- build ------------------------------------------------------------------
def inspect_source(wb):
    print(f"source {SRC}")
    for ws in wb.worksheets:
        s = structure(ws)
        ps = ws.page_setup
        print(f"  [{ws.title}] {ws.dimensions}, widths "
              f"{ {k: v.width for k, v in ws.column_dimensions.items()} }")
        print(f"    merges {sorted(str(m) for m in ws.merged_cells.ranges)}")
        print(f"    page {ps.orientation} paper={ps.paperSize} fitToWidth={ps.fitToWidth} "
              f"fitToHeight={ps.fitToHeight} fitToPage={ws.sheet_properties.pageSetUpPr.fitToPage} "
              f"print area={ws.print_area} titles={ws.print_title_rows}")
        print(f"    {len(s['items'])} items ({sum(s['priced'].values())} priced): "
              f"{', '.join(s['items'])}")
        for sec, (r0, r1, rsub) in s["sections"].items():
            print(f"    section {sec}: items rows {r0}-{r1}, subtotal F{rsub} {ws.cell(rsub, 6).value}")
        lr = s["lot_row"]
        print(f"    LOT subtotal F{lr} {ws.cell(lr, 6).value}; rows {lr + 1}-{ws.max_row} "
              f"(REKAPITULACIJA/NAPOMENA) are dropped")


def site_sheet(wb, src, lot, site, lot1_src, log):
    ws = wb.copy_worksheet(src)
    ws.title = SHEET[lot, site]
    ws.HeaderFooter = copy(src.HeaderFooter)          # "[BH Telecom | Interno]" label
    ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    ws.sheet_view.zoomScaleNormal = src.sheet_view.zoomScaleNormal
    s = structure(ws)
    lot_row = s["lot_row"]
    truncate(ws, lot_row)

    ws["A2"] = f"{ws['A2'].value} — {SITE[site]['title']}"
    ws["A3"] = replace_once(ws["A3"].value, SRC_SITE, JOINT, f"{ws.title} A3")
    wrap_title(ws["A3"])
    ws.row_dimensions[3].height = 30
    if not LOT_RE.match(str(ws.cell(lot_row, 2).value)):
        raise ValueError(f"{ws.title}: LOT subtotal label not in column B")
    ws.cell(lot_row, 2).value = f"UKUPNO {lot} — {SITE[site]['up']} (bez PDV-a):"
    apply_text_edits(ws, BOTH_SITES_EDITS.get(lot, {}), log, "source typo fix", idempotent=True)
    apply_note_edits(ws, BOTH_SITES_NOTE_EDITS.get(lot, {}), log, idempotent=True)

    for sec, (_, _, rsub) in s["sections"].items():      # label clipped in column A
        a, b = ws.cell(rsub, 1), ws.cell(rsub, 2)
        if isinstance(a.value, str) and SUB_RE.match(a.value) and b.value is None:
            b.value, a.value = a.value, None
            b._style = copy(a._style)
            log.append(f"{ws.title}: '{b.value[:10]}...' label moved A{rsub} -> B{rsub}")
    if lot == "LOT 2":
        for col in "EF":
            ws.column_dimensions[col].width = lot1_src.column_dimensions[col].width
    fixed = normalize_item_styles(ws)
    if fixed:
        log.append(f"{ws.title}: borders/number format normalized in {', '.join(fixed)}")
    ws.print_title_rows = PRINT_TITLES
    return ws


def fix_text(text, pairs, where, idempotent=False):
    """Exact (old, new) replacements; each old must occur exactly once. With idempotent=True
    an old that is gone while its new is present counts as already done."""
    done = 0
    for old, new in pairs:
        if idempotent and old not in text and new in text:
            continue
        text = replace_once(text, old, new, where)
        done += 1
    return text, done


def apply_text_edits(ws, edits, log, what="text edit", idempotent=False):
    """Exact (old, new) replacements inside item texts (see fix_text)."""
    items = structure(ws)["items"]
    for n, pairs in edits.items():
        c = ws.cell(items[n], 2)
        c.value, done = fix_text(c.value, pairs, f"{ws.title} {n}", idempotent)
        log.append(f"{ws.title} {n}: {done} {what}(s)"
                   + (f", {len(pairs) - done} already in the source" if done < len(pairs) else ""))


def note_row(ws, start):
    """The one row without item number whose column-B text starts with `start`."""
    rows = [r for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(r, 2).value, str) and ws.cell(r, 2).value.startswith(start)]
    if len(rows) != 1:
        raise ValueError(f"{ws.title}: note {start!r} found {len(rows)} times")
    return rows[0]


def apply_note_edits(ws, edits, log, idempotent=False):
    """Exact (old, new) replacements in a row without item number, found by its text start."""
    for start, pairs in edits.items():
        r = note_row(ws, start)
        c = ws.cell(r, 2)
        c.value, done = fix_text(c.value, pairs, f"{ws.title} B{r}", idempotent)
        log.append(f"{ws.title} B{r} ({start[:22]}...): {done} text edit(s)"
                   + (f", {len(pairs) - done} already in the source" if done < len(pairs) else ""))


def apply_hamzici_lot1(ws, log):
    apply_text_edits(ws, HAMZICI_LOT1_EDITS, log)
    insert_new_items(ws, "LOT 1", log)


def apply_hamzici_lot2(ws, log):
    apply_text_edits(ws, HAMZICI_EDITS, log)
    apply_note_edits(ws, HAMZICI_NOTE_EDITS["LOT 2"], log)
    items = structure(ws)["items"]
    for n, (start, unit, qty, text) in HAMZICI_REPLACE.items():
        r = items[n]
        if not str(ws.cell(r, 2).value).startswith(start):
            raise ValueError(f"{ws.title} {n}: unexpected source text")
        ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value = text, unit, qty
        log.append(f"{n}: replaced ({unit} x {qty})")
    for n, (unit, qty) in HAMZICI_QTY.items():
        r = items[n]
        if ws.cell(r, 3).value != unit:
            raise ValueError(f"{ws.title} {n}: unit {ws.cell(r, 3).value!r} != {unit!r}")
        log.append(f"{n}: qty {ws.cell(r, 4).value} -> {qty} {unit}")
        ws.cell(r, 4).value = qty
    insert_new_items(ws, "LOT 2", log)


def minor(n):
    """Minor number of an item as an integer: "4.10" -> 10, "2.2a" -> 2."""
    return int(re.match(r"\d+", n.split(".")[1]).group(0))


def insert_new_items(ws, which, log):
    """NEW_ITEMS[which]: each row goes right after its anchor item; its number must be free
    and fall numerically between the anchor and the section's next item."""
    for n, after, unit, qty, text in NEW_ITEMS[which]:
        items = structure(ws)["items"]
        sec = n.split(".")[0]
        if n in items or after.split(".")[0] != sec:
            raise ValueError(f"{ws.title}: {n} exists already, or {after} is in another section")
        order = [m for m in items if m.split(".")[0] == sec]
        k = order.index(after)
        nxt = order[k + 1] if k + 1 < len(order) else None
        if not (minor(after) < minor(n) and (nxt is None or minor(n) < minor(nxt))):
            raise ValueError(f"{ws.title}: {n} does not fit between {after} and {nxt}")
        at = items[after] + 1
        insert_row(ws, at)
        for col in range(1, 7):
            ws.cell(at, col)._style = copy(ws.cell(at - 1, col)._style)
        ws.cell(at, 1).value, ws.cell(at, 2).value = n, text
        ws.cell(at, 3).value, ws.cell(at, 4).value = unit, qty
        ws.cell(at, 6).value = line_formula(at)
        # openpyxl moves cells but not formulas: shift_formula_rows() applied Excel's
        # insertion rule above; the three formulas the new row changes are rewritten here.
        s = structure(ws)
        r0, r1, rsub = s["sections"][n.split(".")[0]]
        if not r0 <= at <= r1:
            raise ValueError(f"{ws.title}: {n} (row {at}) fell outside its section {r0}-{r1}")
        before = ws.cell(rsub, 6).value
        ws.cell(rsub, 6).value = f"=SUM(F{r0}:F{rsub - 1})"
        lot = "=" + "+".join(f"F{v[2]}" for v in sorted(s["sections"].values(), key=lambda v: v[2]))
        lot_before = ws.cell(s["lot_row"], 6).value
        ws.cell(s["lot_row"], 6).value = lot
        log.append(f"{n}: new row {at} ({unit} x {qty}); F{rsub} {before} -> {ws.cell(rsub, 6).value}; "
                   f"F{s['lot_row']} {lot_before} -> {lot}")


def style_row(ws, r, src, sr, cols="ABCDEF"):
    for col in cols:
        ws[f"{col}{r}"]._style = copy(src[f"{col}{sr}"]._style)


def build_rekap(wb, lot, src1, src2, log):
    """REKAPITULACIJA of one LOT: its two site totals, discount, 17 % VAT, notes, date and
    signature. Styles and page setup come from the source LOT 2 sheet, as before: portrait
    A4, so the totals, the notes and the signature print on one page for either LOT."""
    src_lot = src1 if lot == "LOT 1" else src2
    ws = wb.create_sheet(REKAP[lot])
    for col in "ABCD":
        ws.column_dimensions[col].width = src2.column_dimensions[col].width
    for col in "EF":
        ws.column_dimensions[col].width = src1.column_dimensions[col].width
    ws.sheet_format = copy(src2.sheet_format)
    ws.sheet_view.zoomScale = src2.sheet_view.zoomScale
    ws.sheet_view.zoomScaleNormal = src2.sheet_view.zoomScaleNormal

    ws["A1"] = src2["A1"].value
    style_row(ws, 1, src2, 1, "A")
    ws["A2"] = f"OBRAZAC ZA CIJENU PONUDE - {lot} - REKAPITULACIJA"
    ws["A3"] = replace_once(src1["A3"].value, SRC_SITE, JOINT, f"{REKAP[lot]} A3")
    for r in (2, 3):
        style_row(ws, r, src2, r, "A")
        ws.merge_cells(f"A{r}:F{r}")
    wrap_title(ws["A3"])
    for r in (1, 2):
        ws.row_dimensions[r].height = src2.row_dimensions[r].height
    ws.row_dimensions[3].height = 30
    for r in (5, 6):
        ws[f"B{r}"] = src2[f"B{r}"].value
        style_row(ws, r, src2, r, "B")
        ws.row_dimensions[r].height = src2.row_dimensions[r].height

    # source styles: 65 = REKAPITULACIJA head, 14 = section head, 66 = value row,
    # 68 = total row (also the grand total: row 63 printed it unstyled), 74/75 = NAPOMENA
    head, sect, val, tot = 65, 14, 66, 68
    r = 8
    style_row(ws, r, src2, head)
    ws[f"B{r}"] = "REKAPITULACIJA"
    r += 1
    style_row(ws, r, src2, sect)
    ws[f"B{r}"] = src_lot["B10"].value
    first = r + 1
    for site in SITE:
        r += 1
        style_row(ws, r, src2, val)
        style_row(ws, r, src2, 16, "B")               # Arial 9 like the item texts
        name = SHEET[lot, site]
        ws[f"B{r}"] = LBL_SITE[lot, site]
        ws[f"F{r}"] = f"='{name}'!F{structure(wb[name])['lot_row']}"
    r_lot, r_pop, r_disc, r_vat, r_grand = r + 1, r + 2, r + 3, r + 4, r + 5
    style_row(ws, r_lot, src2, tot)
    ws[f"B{r_lot}"] = LBL_LOT_TOTAL[lot]
    ws[f"F{r_lot}"] = f"=SUM(F{first}:F{r})"
    style_row(ws, r_pop, src2, tot, "AB")
    style_row(ws, r_pop, src2, val, "CDEF")          # bordered input cell for the discount
    ws[f"B{r_pop}"] = LBL_POPUST
    style_row(ws, r_disc, src2, tot)
    ws[f"B{r_disc}"] = LBL_DISC[lot]
    ws[f"F{r_disc}"] = f'=F{r_lot}*(1-IF(F{r_pop}="",0,F{r_pop}/100))'
    style_row(ws, r_vat, src2, tot)
    ws[f"B{r_vat}"] = LBL_VAT
    ws[f"F{r_vat}"] = f"=F{r_disc}*0.17"
    style_row(ws, r_grand, src2, tot)
    ws[f"B{r_grand}"] = LBL_GRAND[lot]
    ws[f"F{r_grand}"] = f"=F{r_disc}+F{r_vat}"
    for rr in range(8, r_grand + 1):
        ws.row_dimensions[rr].height = 15

    # NAPOMENA - the source notes, checked verbatim, then adapted to two sites and one LOT
    r_note = find_row(src2, "NAPOMENA:")
    src_notes = []
    rr = r_note + 1
    while isinstance(src2.cell(rr, 2).value, str) and src2.cell(rr, 2).value.startswith("- "):
        src_notes.append(src2.cell(rr, 2).value)
        rr += 1
    expected = [s for s, _ in NOTES if s is not None]
    if src_notes != expected:
        raise ValueError(f"source NAPOMENA changed:\n{src_notes}")
    sheets = " i ".join(f"„{n}“" for n in SITE_SHEETS[lot])
    r = r_grand + 2
    style_row(ws, r, src2, r_note, "AB")
    ws[f"B{r}"] = "NAPOMENA:"
    ws.row_dimensions[r].height = 15
    for source, edits in NOTES:
        r += 1
        style_row(ws, r, src2, r_note + 1, "AB")
        if source is None:
            text = edits.format(lot=lot, sheets=sheets)
        else:
            text = source
            for old, new in edits:
                text = replace_once(text, old, new, "NAPOMENA")
        ws[f"B{r}"] = text

    r_date = find_row(src2, "Datum:  ___________________________")
    r += 2
    ws[f"B{r}"] = src2[f"B{r_date}"].value
    style_row(ws, r, src2, r_date, "B")
    r += 2
    ws[f"B{r}"] = src2[f"E{r_date}"].value                # Potpis i pečat Ponuđača
    ws[f"B{r}"]._style = copy(src2[f"E{r_date}"]._style)
    al = copy(ws[f"B{r}"].alignment)
    al.horizontal = "right"
    ws[f"B{r}"].alignment = al
    for rr in (r - 2, r):
        ws.row_dimensions[rr].height = 15

    ws.page_setup.orientation = src2.page_setup.orientation
    ws.page_setup.paperSize = src2.page_setup.paperSize
    ws.page_margins = copy(src2.page_margins)
    ws.print_options = copy(src2.print_options)
    ws.HeaderFooter = copy(src2.HeaderFooter)
    print_setup(ws, r)
    log.append(f"{REKAP[lot]}: totals rows 10-{r_grand}, {len(NOTES)} notes, signature row {r}")
    return ws


def validate(wb, lot, src_formulas):
    """Structural self-check before saving."""
    for name in SITE_SHEETS[lot]:
        ws = wb[name]
        s = structure(ws)
        bad = []
        for n, r in s["items"].items():
            if s["priced"][n] and ws.cell(r, 6).value != line_formula(r):
                bad.append(f"{n} F{r}")
            if ws.cell(r, 5).value not in (None, ""):
                bad.append(f"{n} has a unit price")
        order = sorted(s["sections"].items(), key=lambda kv: kv[1][2])
        for sec, (r0, r1, rsub) in order:
            if ws.cell(rsub, 6).value != f"=SUM(F{r0}:F{rsub - 1})" or r1 > rsub - 1:
                bad.append(f"section {sec} F{rsub}={ws.cell(rsub, 6).value}")
        want = "=" + "+".join(f"F{v[2]}" for _, v in order)
        if ws.cell(s["lot_row"], 6).value != want:
            bad.append(f"LOT F{s['lot_row']}={ws.cell(s['lot_row'], 6).value} != {want}")
        if ws.max_row != s["lot_row"]:
            bad.append(f"rows below the LOT subtotal (max_row {ws.max_row})")
        merges = sorted(str(m) for m in ws.merged_cells.ranges)
        if merges != ["A2:F2", "A3:F3"]:
            bad.append(f"merges {merges}")
        if name in src_formulas:                       # sheets without inserted rows
            now = {c.coordinate: c.value for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str) and c.value.startswith("=")}
            if now != src_formulas[name]:
                diff = {k for k in set(now) | set(src_formulas[name])
                        if now.get(k) != src_formulas[name].get(k)}
                bad.append(f"formulas differ from source at {sorted(diff)}")
        if bad:
            raise AssertionError(f"{name}: {bad}")


def build(lot):
    """One LOT's price form: its two site sheets and its REKAPITULACIJA."""
    wb = openpyxl.load_workbook(SRC)
    src1, src2 = wb["LOT 1"], wb["LOT 2"]
    src = src1 if lot == "LOT 1" else src2
    last = structure(src)["lot_row"]
    snap = {c.coordinate: c.value for row in src.iter_rows(max_row=last) for c in row
            if isinstance(c.value, str) and c.value.startswith("=")}

    fixes, changes = [], []
    for site in SITE:
        site_sheet(wb, src, lot, site, src1, fixes)
    apply = apply_hamzici_lot1 if lot == "LOT 1" else apply_hamzici_lot2
    apply(wb[SHEET[lot, "hamzici"]], changes)
    build_rekap(wb, lot, src1, src2, fixes)

    heights = {}
    for name in SITE_SHEETS[lot]:
        ws = wb[name]
        heights[name] = autofit_rows(ws)
        print_setup(ws, ws.max_row)
    heights[REKAP[lot]] = autofit_rows(wb[REKAP[lot]])

    wb.remove(src1)
    wb.remove(src2)
    wb._sheets = [wb[n] for n in SHEET_ORDER[lot]]
    for i, ws in enumerate(wb.worksheets):
        ws.sheet_view.tabSelected = i == 0
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True

    validate(wb, lot, {SHEET[lot, "sjednica"]: snap})
    out = OUT[lot]
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)

    print(f"\nwritten {out}")
    for ws in wb.worksheets:
        extra = ""
        if ws.title in SITE_SHEETS[lot]:
            s = structure(ws)
            extra = (f", {len(s['items'])} items, LOT subtotal F{s['lot_row']} = "
                     f"{ws.cell(s['lot_row'], 6).value}")
        print(f"  {ws.title}: {ws.max_row} rows, print area {ws.print_area}, "
              f"{heights[ws.title]} rows set to auto-fit{extra}")
    print(f"{lot} Hamzići changes:")
    for c in changes:
        print("  " + c)
    print("fixes / notes:")
    for f in fixes:
        print("  " + f)
    props = getattr(wb, "custom_doc_props", None)
    if props is not None:
        print(f"custom document properties kept: {[p.name for p in props.props]}")


def main():
    inspect_source(openpyxl.load_workbook(SRC))
    for lot in LOTS:
        build(lot)


if __name__ == "__main__":
    main()
