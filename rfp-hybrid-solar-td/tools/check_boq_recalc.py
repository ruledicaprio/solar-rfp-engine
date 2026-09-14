# -*- coding: utf-8 -*-
"""Recalculation check for the joint price forms (Prilog II, one workbook per LOT).

Copies the two workbooks written by joint_boq.py to a scratch folder, puts a
unit price of 100 KM into every priced item row, lets LibreOffice (headless,
private profile) recalculate them, and checks the recalculated values against
sums computed here in Python:

  * every line total = quantity x 100;
  * every section subtotal, every LOT subtotal and every total of the LOT's own
    REKAPITULACIJA (discount, 17 % VAT) equals the Python sum (run twice: no
    discount, and a 5 % discount);
  * no #REF!/#VALUE!/#NAME?/... anywhere, before or after recalculation;
  * the delivered files carry no prices and no simulation/energy figures.

It also prints a COVERAGE diff: every Sjednica item number must exist in the
Hamzići sheet of the same LOT, and every Hamzići-only change is listed.

    python check_boq_recalc.py [--xlsx LOT1_FILE LOT2_FILE] [--scratch DIR]

Exit code 0 = all checks passed.
"""
import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import joint_boq as jb  # noqa: E402  (sheet names, labels, structure helpers)
import paths  # noqa: E402

PRICE = 100.0
DISCOUNT = 5.0
VAT = 0.17
TOL = 1e-6
ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
# Figures that only a simulation/energy balance produces - none may appear in the BoQ.
ENERGY_RE = re.compile(r"kWh|MWh|h/god|l/god|sati rada godi|solarn\w* udio|solar fraction|pvsim|PVGIS|"
                       r"potrošnj\w* goriva", re.I)
# Mentions that are worth a look but come from the source text (reported, not failed).
ENERGY_INFO_RE = re.compile(r"zračenj|prinos|irradi", re.I)

FAILS = []


def fail(msg):
    FAILS.append(msg)
    print("  FAIL:", msg)


def ok(msg):
    print("  ok  :", msg)


def soffice():
    exe = os.environ.get("SOFFICE") or shutil.which("soffice") or shutil.which("soffice.exe")
    if not exe:
        exe = r"C:\Program Files\LibreOffice\program\soffice.exe"
    if not os.path.exists(exe):
        sys.exit("LibreOffice not found (set SOFFICE)")
    return exe


def ensure_recalc_profile(profile):
    """Make the private LibreOffice profile recalculate OOXML formulas on load."""
    user = os.path.join(profile, "user")
    os.makedirs(user, exist_ok=True)
    reg = os.path.join(user, "registrymodifications.xcu")
    items = ('<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
             '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n'
             '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
             '<prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n')
    if os.path.exists(reg):
        with open(reg, encoding="utf-8") as f:
            xml = f.read()
        if "OOXMLRecalcMode" in xml:
            xml = re.sub(r'(<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>)\d(</value>)', r"\g<1>0\g<2>", xml)
        else:
            xml = xml.replace("</oor:items>", items + "</oor:items>")
    else:
        xml = ('<?xml version="1.0" encoding="UTF-8"?>\n'
               '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
               'xmlns:xs="http://www.w3.org/2001/XMLSchema" '
               'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n' + items + "</oor:items>\n")
    with open(reg, "w", encoding="utf-8") as f:
        f.write(xml)


def recalc(files, outdir, profile):
    ensure_recalc_profile(profile)
    os.makedirs(outdir, exist_ok=True)
    uri = "file:///" + os.path.abspath(profile).replace("\\", "/").lstrip("/")
    cmd = [soffice(), f"-env:UserInstallation={uri}", "--headless", "--norestore",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outdir] + list(files)
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    out = [os.path.join(outdir, os.path.basename(f)) for f in files]
    missing = [o for o in out if not os.path.exists(o)]
    if res.returncode != 0 or missing:
        sys.exit(f"LibreOffice conversion failed (rc={res.returncode}): {res.stdout}\n{res.stderr}")
    return out


def scan_errors(wb, label):
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and any(e in v for e in ERRORS):
                    bad.append(f"{ws.title}!{c.coordinate}={v[:40]!r}")
    if bad:
        fail(f"{label}: error values {bad[:10]}{' ...' if len(bad) > 10 else ''}")
    else:
        ok(f"{label}: no #REF!/#VALUE!/#NAME?/Err: values")


def check_delivered(wb, lot):
    """The file as delivered: formulas only, no prices, no energy figures."""
    print(f"\n[1] Delivered file, {lot}")
    names = wb.sheetnames
    if names != jb.SHEET_ORDER[lot]:
        fail(f"{lot}: sheet order {names} != {jb.SHEET_ORDER[lot]}")
    else:
        ok(f"sheets {names}")
    scan_errors(wb, "formulas")
    info = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    if ENERGY_RE.search(c.value):
                        fail(f"energy/simulation figure in {ws.title}!{c.coordinate}: "
                             f"{ENERGY_RE.search(c.value).group(0)!r}")
                    m = ENERGY_INFO_RE.search(c.value)
                    if m:
                        info.append(f"{ws.title}!{c.coordinate} ({m.group(0)})")
    if info:
        print("  info: source text mentions irradiation/yield in", ", ".join(info))
    for name in jb.SITE_SHEETS[lot]:
        ws = wb[name]
        s = jb.structure(ws)
        priced = [n for n, r in s["items"].items() if s["priced"][n]]
        with_price = [n for n in priced if ws.cell(s["items"][n], 5).value not in (None, "")]
        consts = [n for n in priced if not str(ws.cell(s["items"][n], 6).value or "").startswith("=")]
        if with_price:
            fail(f"{name}: unit prices present in {with_price}")
        if consts:
            fail(f"{name}: line total is not a formula in {consts}")
        own = [n for n in priced
               if ws.cell(s["items"][n], 6).value != jb.line_formula(s["items"][n])]
        if own:
            fail(f"{name}: line formula does not reference its own row in {own}")
        if not (with_price or consts or own):
            ok(f"{name}: {len(priced)} priced items, E empty, F = own-row formula; "
               f"merges {sorted(str(m) for m in ws.merged_cells.ranges)}; "
               f"print area {ws.print_area}; titles {ws.print_title_rows}; "
               f"fitToPage {ws.sheet_properties.pageSetUpPr.fitToPage} "
               f"({ws.page_setup.fitToWidth}x{ws.page_setup.fitToHeight})")


def price_copy(src, dst, discount, lot):
    wb = openpyxl.load_workbook(src)
    n = 0
    for name in jb.SITE_SHEETS[lot]:
        ws = wb[name]
        s = jb.structure(ws)
        for item, r in s["items"].items():
            if s["priced"][item]:
                ws.cell(r, 5).value = PRICE
                n += 1
    rk = wb[jb.REKAP[lot]]
    rk.cell(jb.find_row(rk, jb.LBL_POPUST), 6).value = discount
    wb.save(dst)
    return n


def expected(wb_formulas, lot):
    """Python-side totals from quantities (the unrecalculated workbook)."""
    exp = {}
    for name in jb.SITE_SHEETS[lot]:
        ws = wb_formulas[name]
        s = jb.structure(ws)
        lines = {r: float(ws.cell(r, 4).value) * PRICE
                 for item, r in s["items"].items() if s["priced"][item]}
        secs = {}
        for sec, (r0, r1, rsub) in s["sections"].items():
            secs[sec] = (rsub, sum(v for r, v in lines.items() if r0 <= r <= r1))
        exp[name] = {"lines": lines, "sections": secs,
                     "lot": (s["lot_row"], sum(v for _, v in secs.values()))}
    return exp


def check_recalc(path, exp, discount, lot):
    print(f"\n[2] {lot} recalculated by LibreOffice, unit price {PRICE:g}, discount {discount:g} %")
    wb = openpyxl.load_workbook(path, data_only=True)
    scan_errors(wb, f"{lot} values")
    for name in jb.SITE_SHEETS[lot]:
        ws = wb[name]
        e = exp[name]
        bad = [(r, ws.cell(r, 6).value, v) for r, v in e["lines"].items()
               if not isinstance(ws.cell(r, 6).value, (int, float)) or abs(ws.cell(r, 6).value - v) > TOL]
        if bad:
            fail(f"{name}: line totals != qty x {PRICE:g} at rows {bad[:5]}")
        else:
            ok(f"{name}: {len(e['lines'])} line totals = qty x {PRICE:g}")
        for sec, (rsub, v) in e["sections"].items():
            got = ws.cell(rsub, 6).value
            if not isinstance(got, (int, float)) or abs(got - v) > TOL:
                fail(f"{name}: section {sec} subtotal F{rsub}={got} != {v:.2f}")
            else:
                ok(f"{name}: section {sec} subtotal F{rsub} = {got:,.2f}")
        rl, v = e["lot"]
        got = ws.cell(rl, 6).value
        if not isinstance(got, (int, float)) or abs(got - v) > TOL:
            fail(f"{name}: LOT subtotal F{rl}={got} != {v:.2f}")
        else:
            ok(f"{name}: LOT subtotal F{rl} = {got:,.2f}")
    rk = wb[jb.REKAP[lot]]
    site = {s: exp[jb.SHEET[lot, s]]["lot"][1] for s in jb.SITE}
    tot = sum(site.values())
    disc = tot * (1 - discount / 100.0)
    want = {jb.LBL_SITE[lot, s]: v for s, v in site.items()}
    want.update({
        jb.LBL_LOT_TOTAL[lot]: tot,
        jb.LBL_POPUST: discount,
        jb.LBL_DISC[lot]: disc,
        jb.LBL_VAT: disc * VAT,
        jb.LBL_GRAND[lot]: disc * (1 + VAT),
    })
    for label, v in want.items():
        r = jb.find_row(rk, label)
        got = rk.cell(r, 6).value
        if label == jb.LBL_POPUST and not discount:      # left empty = no discount
            if got not in (None, ""):
                fail(f"{rk.title} F{r} '{label}' should be empty, is {got!r}")
            else:
                ok(f"{rk.title} F{r} {label} empty (no discount)")
            continue
        if not isinstance(got, (int, float)) or abs(got - v) > TOL:
            fail(f"{rk.title} F{r} '{label}' = {got} != {v:.2f}")
        else:
            ok(f"{rk.title} F{r} {label} {got:,.2f}")


def coverage(wbs):
    print("\n[3] COVERAGE (Sjednica -> Hamzići)")
    for lot in jb.LOTS:
        wb = wbs[lot]
        sj, hz = wb[jb.SHEET[lot, "sjednica"]], wb[jb.SHEET[lot, "hamzici"]]
        a, b = jb.structure(sj)["items"], jb.structure(hz)["items"]
        missing = [n for n in a if n not in b]
        extra = [n for n in b if n not in a]
        if missing:
            fail(f"{lot}: Sjednica items missing in Hamzići: {missing}")
        else:
            ok(f"{lot}: all {len(a)} Sjednica items present in Hamzići ({len(b)} items)")
        changes = []
        for n in a:
            if n not in b:
                continue
            ra, rb = a[n], b[n]
            diff = []
            if sj.cell(ra, 2).value != hz.cell(rb, 2).value:
                diff.append("text")
            if sj.cell(ra, 3).value != hz.cell(rb, 3).value:
                diff.append(f"unit {sj.cell(ra, 3).value}->{hz.cell(rb, 3).value}")
            if sj.cell(ra, 4).value != hz.cell(rb, 4).value:
                diff.append(f"qty {sj.cell(ra, 4).value}->{hz.cell(rb, 4).value}")
            if diff:
                changes.append(f"{n} ({', '.join(diff)})")
        for n in extra:
            changes.append(f"{n} (NEW: {hz.cell(b[n], 3).value} x {hz.cell(b[n], 4).value}: "
                           f"{str(hz.cell(b[n], 2).value)[:60]}...)")
        print(f"  {lot} Hamzići-only changes: {', '.join(changes) if changes else 'none'}")
        expected_changes = jb.EXPECTED_CHANGES[lot]
        got = {c.split(" ")[0] for c in changes}
        if got != set(expected_changes):
            fail(f"{lot}: changed items {sorted(got)} != intended {sorted(expected_changes)}")
        else:
            ok(f"{lot}: changes are exactly the intended ones")


# Sjednica-only facts (and superseded Hamzići values) that must not survive in a Hamzići
# sheet - including the Sjednica drawing numbers, now that the H-01..H-05 map is in.
SJEDNICA_FACTS = re.compile(r"1076|42,94|16,00 × 9,40|h=2,10|h=38|planinsk|11,6 kW|≈82 %|93 %|"
                            r"trase od 25 m|dužini 25 m|900 × 2000|jugoistočne strane ograde|"
                            r"1,94 m|4\.4\.2\.3|720 mm|1155 mm|CENTRIRANO|"
                            r"JUŽNI ugao kontejnera \(uz JZ|cca 0,40 m od ograde|0,83 m|"
                            r"0,60 × 0,25|≈1,9 kPa|dužine do 15 m|\b[MSE]-0\d\b|"
                            # superseded Hamzići layout: Stulz and discharge on the east wall
                            r"ISTOČNE strane je hladnjak|južnom kraju istočnog zida|"
                            r"uz ISTOČNI zid kontejnera, sjeverno od agregata|"
                            r"sa istočnog zida kontejnera|Stulz \(Tačka 5\.15\) na ISTOČNOM|"
                            r"ugrađuje u ISTOČNI zid kontejnera, na mjestu|"
                            # superseded Hamzići layout of the plan-frame round (before the
                            # 45° orientation, 11.09.2026): walls named as if door = north
                            r"sredini JUŽNOG zida|u ZAPADNI zid|JUGOZAPADNOM uglu|kroz ISTOČNI "
                            r"zid|sjevernom pojasu|na JUŽNOM zidu|pojasu južno od ploče|1,74 m")


def check_sjednica(wbs):
    print("\n[4] Sjednica sheets vs source; Sjednica facts left in Hamzići sheets")
    src = openpyxl.load_workbook(jb.SRC)
    for lot in ("LOT 1", "LOT 2"):
        wb = wbs[lot]
        sj, s = wb[jb.SHEET[lot, "sjednica"]], src[lot]
        a, b = jb.structure(sj)["items"], jb.structure(s)["items"]
        # items added to BOTH site sheets are not in the source: expect them at their
        # anchor position, and check them against what joint_boq declares, not the source
        added = {n: (unit, qty, text)
                 for n, after, unit, qty, text in jb.BOTH_SITES_NEW_ITEMS.get(lot, [])}
        want_list = list(b)
        for n, after, *_ in jb.BOTH_SITES_NEW_ITEMS.get(lot, []):
            want_list.insert(want_list.index(after) + 1, n)
        if list(a) != want_list:
            fail(f"{sj.title}: item list differs from the source" +
                 (f" + both-sites items {sorted(added)}" if added else ""))
        else:
            # expected = the source text, plus any typo fix the source does not carry yet
            bad, fixed = [], []
            for n in a:
                if n in added:
                    unit, qty, text = added[n]
                    if (sj.cell(a[n], 2).value, sj.cell(a[n], 3).value,
                            sj.cell(a[n], 4).value) != (text, unit, qty):
                        bad.append(n)
                    continue
                want = s.cell(b[n], 2).value
                pairs = jb.BOTH_SITES_EDITS.get(lot, {}).get(n)
                if pairs:
                    want, done = jb.fix_text(want, pairs, n, idempotent=True)
                    fixed += [n] if done else []
                if sj.cell(a[n], 2).value != want or any(
                        sj.cell(a[n], c).value != s.cell(b[n], c).value for c in (3, 4)):
                    bad.append(n)
            for start, pairs in jb.BOTH_SITES_NOTE_EDITS.get(lot, {}).items():
                want, done = jb.fix_text(s.cell(jb.note_row(s, start), 2).value, pairs, start,
                                         idempotent=True)
                fixed += [start[:24]] if done else []
                if sj.cell(jb.note_row(sj, start), 2).value != want:
                    bad.append(start[:24])
            if bad:
                fail(f"{sj.title}: differs from the source in {bad}")
            else:
                ok(f"{sj.title}: equal to the source ({len(b)} items)"
                   + (f" + {len(added)} both-sites item(s) {sorted(added)}" if added else "")
                   + (f"; typo fixes the source lacks: {fixed}" if fixed
                      else "; every typo fix is already in the source"))
        hz = wb[jb.SHEET[lot, "hamzici"]]
        hits = []
        for r in range(1, hz.max_row + 1):
            v = hz.cell(r, 2).value
            m = SJEDNICA_FACTS.search(v) if isinstance(v, str) else None
            if m:
                hits.append(f"{hz.cell(r, 1).value or 'B%d' % r}: {m.group(0)!r}")
        if hits:
            fail(f"{hz.title}: Sjednica-only facts left in {hits}")
        else:
            ok(f"{hz.title}: no Sjednica-only facts (1076 m, 42,94° N, 38 m mast, 16,00 × 9,40, ...)")


# Always-on loads on −48 V DC (Naručilac, 11.09.2026, both sites): source items 5.16-5.18.
DC_ITEMS = ("5.16", "5.17", "5.18")
DC_COVERAGE = ["−48 V", "ICC360-HA1-C1", "2 × 6 mm²", "48 V DC", "EC ventilator", "EN 54-4",
               "DC/DC", "foto-senzor", "Prilog I, Tačka 4.5", "prekidačem uz ulazna vrata",
               "PREDGRIJAČ RASHLADNE TEČNOSTI", "predgrijač rashladne tečnosti agregata",
               "Prilog I, Tačka 4.1", '1-„agregat", 0-„isključeno", 2-„rezerva"',
               "izlaz kontrolera (relej)", "položaj 2, npr. mobilni agregat",
               "zimskim dizel gorivom prema EN 590", "zadano −10 °C"]
# an always-on load left on the GRO (dead while the genset is stopped), the emergency
# luminaire the DC light replaces, the dropped 230 V / space / tank heaters, the old switch
# labels and switching sentence, the old 10 °C room alarm
ON_GRO = re.compile(r"5 A / 230 V|ZASEBAN NADZIRANI|nadzirani strujni krug postoje|protupani|"
                    r"RASHLADNE TEČNOSTI 230 V|grijač prostora|„hibridni sistem|"
                    r"strujni krug za grijač motora|između hibridnog sistema napajanja i agregata|"
                    r"grijanje/zaštita spremnika|t<10 °C")


def check_dc(wb):
    print("\n[5] Always-on loads on −48 V DC (5.16-5.18, 4.8, 5.10, 3.1)")
    for site in jb.SITE:
        ws = wb[jb.SHEET["LOT 2", site]]
        s = jb.structure(ws)
        items = s["items"]
        missing = [n for n in DC_ITEMS if n not in items]
        if missing:
            fail(f"{ws.title}: DC items missing {missing}")
            continue
        r0, r1, rsub = s["sections"]["5"]
        bad = [n for n in DC_ITEMS
               if not (s["priced"][n] and r0 <= items[n] <= r1
                       and (ws.cell(items[n], 3).value, ws.cell(items[n], 4).value) == ("kpl", 1))]
        sec5 = [n for n in items if n.split(".")[0] == "5"]
        order_ok = all(jb.minor(p) < jb.minor(q) for p, q in zip(sec5, sec5[1:]))
        text = "\n".join(str(ws.cell(r, 2).value) for r in range(1, ws.max_row + 1)
                         if ws.cell(r, 2).value)
        cov = [t for t in DC_COVERAGE if t not in text]
        gro = sorted(set(ON_GRO.findall(text)))
        if bad or not order_ok or cov or gro:
            fail(f"{ws.title}: DC items not priced/kpl×1/inside UKUPNO 5 {bad}; section 5 order "
                 f"{'ok' if order_ok else sec5}; coverage missing {cov}; on GRO {gro}")
        else:
            ok(f"{ws.title}: {', '.join(DC_ITEMS)} priced kpl × 1 inside UKUPNO 5 (F{rsub} = "
               f"{ws.cell(rsub, 6).value}); section 5 in order ({' '.join(sec5[-5:])}); "
               f"{len(DC_COVERAGE)} DC terms present; no always-on load on the GRO")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", nargs=2, metavar=("LOT1", "LOT2"),
                    default=[paths.PRILOG2_LOT[lot] for lot in jb.LOTS])
    ap.add_argument("--scratch", default=os.environ.get("BOQ_SCRATCH")
                    or os.path.join(tempfile.gettempdir(), "boq_check"))
    args = ap.parse_args()
    scratch = os.path.abspath(args.scratch)
    os.makedirs(scratch, exist_ok=True)
    print(f"scratch  {scratch}")

    wbs, exp, priced = {}, {}, []
    for lot, path in zip(jb.LOTS, args.xlsx):
        tag = lot.replace(" ", "").lower()
        src = os.path.join(scratch, f"joint_boq_{tag}.xlsx")
        shutil.copyfile(path, src)
        print(f"\nchecking {path}")
        wbs[lot] = openpyxl.load_workbook(src)
        check_delivered(wbs[lot], lot)
        exp[lot] = expected(wbs[lot], lot)
        p0 = os.path.join(scratch, f"priced_{tag}_nodisc.xlsx")
        p5 = os.path.join(scratch, f"priced_{tag}_disc5.xlsx")
        n = price_copy(src, p0, None, lot)
        price_copy(src, p5, DISCOUNT, lot)
        print(f"\n{lot}: unit price {PRICE:g} written into {n} priced item rows")
        priced += [(lot, p0, 0.0), (lot, p5, DISCOUNT)]
    outs = recalc([p for _, p, _ in priced], os.path.join(scratch, "recalc"), os.path.join(scratch, "lo"))
    for (lot, _, discount), out in zip(priced, outs):
        check_recalc(out, exp[lot], discount, lot)
    coverage(wbs)
    check_sjednica(wbs)
    check_dc(wbs["LOT 2"])

    print()
    if FAILS:
        print(f"FAILED: {len(FAILS)} check(s)")
        sys.exit(1)
    print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
