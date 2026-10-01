# -*- coding: utf-8 -*-
"""
Prezentacija planirane nabavke (BH Telecom predložak, 4:3).

    python tools/build_prezentacija.py <predlozak.pptx> [izlaz.pptx]

Predložak je BH Telecom prezentacija planirane nabavke (npr. "Zamjena sistema
klimatizacije TKC Tuzla 4 sprat.pptx"): zadržavaju se master, naslovni slajd i
slajd "HVALA NA PAŽNJI!", ostali slajdovi se brišu, a novi se crtaju na
rasporedu "Title and Content" prvog mastera (narandžasta traka, bh logo,
oznaka "BH Telecom | Interno"). Brojke su iz TD, Priloga I/II i pvsim.
"""
import copy
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
JOINT = os.path.dirname(HERE)
ROOT = os.path.dirname(JOINT)
P3 = os.path.join(JOINT, "TD-OUTPUT", "3.2 Prilog III TD - Situacije.pdf")

FONT = "Bookman Old Style"
ORANGE = RGBColor(0xEF, 0x88, 0x12)
DARK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x5F, 0x5F, 0x5F)
TINT = RGBColor(0xFD, 0xF1, 0xE4)
LIGHT = RGBColor(0xF4, 0xF4, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BOTTOM = 5.85         # traka mastera: ≈6,4" desno, val do ≈5,7" uz lijevu ivicu


# --------------------------------------------------------------------------
# osnovni elementi
def run(p, text, size=13, bold=False, color=DARK, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color
    return r


def bullet(p, char="•", indent=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bf = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    bc = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(bf)
    pPr.append(bc)


def box(slide, x, y, w, h, paras, size=13, anchor=MSO_ANCHOR.TOP, margin=0.0,
        align=PP_ALIGN.LEFT, space=4):
    """paras: list of str | (str, dict) | list of runs [(text, dict), ...]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space)
        if isinstance(para, str):
            para = [(para, {})]
        elif isinstance(para, tuple):
            para = [para]
        opts0 = para[0][1]
        if opts0.get("bullet"):
            bullet(p, indent=opts0.get("indent", 0.22))
        for text, o in para:
            run(p, text, size=o.get("size", size), bold=o.get("bold", False),
                color=o.get("color", DARK), italic=o.get("italic", False))
    return tb


def title(slide, text, sub=None):
    box(slide, 0.4, 0.32, 7.7, 0.6, [(text, {"bold": True, "size": 24})])
    if sub:
        box(slide, 0.4, 0.86, 8.9, 0.35, [(sub, {"size": 12, "color": GREY,
                                                 "italic": True})])


def shape(slide, kind, x, y, w, h, fill=TINT, line=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, fill=TINT):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    s.adjustments[0] = 0.06
    return s


def badge(slide, x, y, d, text, size=13, fill=ORANGE):
    s = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    tf = s.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, text, size=size, bold=True, color=WHITE)
    return s


def arrow(slide, x, y, w, h):
    shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, ORANGE)


def picture(slide, path, x, y, w=None, h=None):
    """Smjesti sliku u okvir (w, h) bez izobličenja; vraća stvarni okvir."""
    iw, ih = Image.open(path).size
    if w and h:
        if iw / ih > w / h:
            h2 = w * ih / iw
            y, h = y + (h - h2) / 2, h2
        else:
            w2 = h * iw / ih
            x, w = x + (w - w2) / 2, w2
    elif w:
        h = w * ih / iw
    else:
        w = h * iw / ih
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return x, y, w, h


def table(slide, x, y, w, rows, widths, size=11, head_fill=ORANGE, row_h=0.3,
          bold_last=False, align=None, head_h=None):
    n, m = len(rows), len(rows[0])
    gt = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w),
                                Inches(row_h * n))
    tbl = gt.table
    # bez ugrađenog stila: bijelo, tanke sive linije
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(head_h if (i == 0 and head_h) else row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.01)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            head = i == 0
            last = bold_last and i == n - 1
            c.fill.fore_color.rgb = head_fill if head else (TINT if last else WHITE)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (align[j] if align else PP_ALIGN.LEFT)
            run(p, str(val), size=size, bold=head or last,
                color=WHITE if head else DARK)
            _borders(c)
    return tbl


def _borders(cell, color="BFBFBF", w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.makeelement(qn(tag), {"w": str(w)})
        sf = ln.makeelement(qn("a:solidFill"), {})
        clr = sf.makeelement(qn("a:srgbClr"), {"val": color})
        sf.append(clr)
        ln.append(sf)
        tcPr.append(ln)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# --------------------------------------------------------------------------
def prep_template(src):
    prs = Presentation(src)
    sld = prs.slides._sldIdLst
    ids = list(sld)
    keep = {0, len(ids) - 1}            # naslovni i "HVALA NA PAŽNJI!"
    for i, sid in enumerate(ids):
        if i not in keep:
            prs.part.drop_rel(sid.get(qn("r:id")))
            sld.remove(sid)
    # novi slajd dobija ime slideN po broju slajdova; zadržani "HVALA" (slide6)
    # bi se sudario sa novim, pa se privremeno sklanja na visok broj
    prs.slides[1].part.partname = PackURI("/ppt/slides/slide999.xml")
    layout = next(l for l in prs.slide_masters[0].slide_layouts
                  if l.name == "Title and Content")
    return prs, layout


def renumber(prs):
    for i, slide in enumerate(prs.slides, 1):
        slide.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")


def new_slide(prs, layout):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def move_last_to_end(prs):
    """Slajd 'HVALA' je drugi u listi; premjesti ga na kraj."""
    sld = prs.slides._sldIdLst
    thanks = list(sld)[1]
    sld.remove(thanks)
    sld.append(thanks)


def edit_title_slide(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("ZAMJENA"):
            _set_lines(sh, ["AUTONOMNI HIBRIDNI SISTEMI NAPAJANJA",
                            "BS SJEDNICA (BILEĆA) I BS HAMZIĆI (ČITLUK)",
                            "",
                            "- prezentacija planirane nabavke (LOT 1 i LOT 2)"],
                       sizes=[24, 20, 12, 16], bold=[True, True, False, False])
        elif "godine" in t:
            _set_lines(sh, ["Oktobar 2026. godine"])
            sh.left, sh.width = Inches(2.5), Inches(5.0)


def _set_lines(sh, lines, sizes=None, bold=None):
    """Zamijeni tekst zadržavajući formatiranje prvog runa svakog paragrafa."""
    tf = sh.text_frame
    paras = list(tf.paragraphs)
    proto = copy.deepcopy(paras[0]._p)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    p0 = paras[0]._p
    for k, line in enumerate(lines):
        p = p0 if k == 0 else copy.deepcopy(proto)
        if k:
            p0.getparent().append(p)
        rs = p.findall(qn("a:r"))
        for r in rs[1:]:
            p.remove(r)
        rs[0].find(qn("a:t")).text = line
        rPr = rs[0].find(qn("a:rPr"))
        if rPr is not None and sizes:
            rPr.set("sz", str(sizes[k] * 100))
        if rPr is not None and bold:
            rPr.set("b", "1" if bold[k] else "0")
        for br in p.findall(qn("a:br")):
            p.remove(br)


# --------------------------------------------------------------------------
def s_investicija(prs, L, img):
    s = new_slide(prs, L)
    title(s, "Opis investicije i postojeće stanje")
    B = {"bullet": True}
    box(s, 0.4, 1.05, 5.4, 3.3, [
        ("BS Sjednica (Bileća) i BS Hamzići (Čitluk) nisu priključene na "
         "elektroenergetsku mrežu; bez napajanja se ne mogu pustiti u rad.", B),
        ("Rješenje je autonomni hibridni sistem na −48 V DC: FN polje, LFP "
         "baterije i dizel agregat kao rezerva.", B),
        ("Huawei oprema (FN moduli, ICC360, baterije, MTS) je posebna nabavka; "
         "ova nabavka obuhvata infrastrukturu i instalaciju.", B),
        ("Druga iteracija, nakon lokacije MILNIŠTE_MIKRO (Glamoč).", B),
    ], size=12, space=7)
    for k, (key, cap) in enumerate((("sjednica", "BS Sjednica"),
                                    ("hamzici", "BS Hamzići"))):
        x, y, w, h = picture(s, img[key], 6.0 + k * 1.85, 1.05, 1.75, 2.95)
        box(s, x, y + h + 0.03, w, 0.28, [(cap, {"size": 9.5, "color": GREY,
                                               "italic": True})],
            align=PP_ALIGN.CENTER)
    rows = [["Lokacija", "Nadm. visina", "Postojeće", "Napomena"],
            ["BS Sjednica, Bileća (RS)", "1076 m", "ploča 5,40 × 5,40 m, stub 38 m, "
             "kontejner K2", "ormar ICC360 uz kontejner"],
            ["BS Hamzići, Čitluk (FBiH)", "493 m", "ploča 5,40 × 5,40 m, stub 32 m, "
             "kontejner K2", "demontaža klime Stulz; makadam ≈800 m"]]
    table(s, 0.4, 4.4, 9.2, rows, [2.35, 1.15, 3.05, 2.65], size=10, row_h=0.45,
          head_h=0.32)
    notes(s, "Bez napajanja nema puštanja u rad; izostaje pokrivanje i RR čvorište "
             "Sjednica. Oba objekta imaju otežan prilaz, naročito nakon padavina.")


def s_rjesenje(prs, L):
    s = new_slide(prs, L)
    title(s, "Tehničko rješenje i predmet nabavke")
    blocks = [(0.4, "FN polje", "12 × 585 Wp\n7,02 kWp"),
              (2.75, "Solarni moduli", "PVDB\n2 × S4875G3"),
              (5.1, "Huawei ICC360", "−48 V DC, SMU\nLFP 48,6 kWh"),
              (7.45, "TK oprema", "≈1,18 kW\nnazivno")]
    for x, h1, h2 in blocks:
        card(s, x, 1.0, 2.15, 1.0)
        box(s, x + 0.08, 1.05, 1.99, 0.9,
            [(h1, {"bold": True, "size": 11, "color": ORANGE}), (h2, {"size": 10})],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=1)
    for x in (2.57, 4.92, 7.27):
        arrow(s, x, 1.37, 0.17, 0.26)
    shape(s, MSO_SHAPE.UP_ARROW, 6.05, 2.03, 0.26, 0.24, ORANGE)
    card(s, 5.1, 2.3, 2.15, 0.62, LIGHT)
    box(s, 5.15, 2.3, 2.05, 0.62, [("Rezerva: DEA 18 kVA", {"bold": True, "size": 10.5,
                                                          "color": ORANGE}),
                                   ("→ GRO → 3 × R4875", {"size": 9.5})],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0)
    box(s, 0.4, 2.2, 4.55, 0.75, [
        [("Očekivano godišnje: ", {"bold": True, "size": 10.5}),
         ("solarni udio 75–77 %, agregat ≈270–300 h, gorivo ≈900–990 l "
          "(pvsim, 2005–2023).", {"size": 10.5})]], space=0)
    lots = [("LOT 1", "Nosači FN panela — 8 kpl (4 po lokaciji)", [
                "zemljani radovi, temeljne trake C30/37",
                "4 nosača × 3 modula, 45°, azimut 225°",
                "hemijski ankeri, uzemljenje nosača",
                "statički proračun — uslov za početak radova"]),
            ("LOT 2", "Agregat i hibridni sistem — 2 kpl", [
                "DEA 18 kVA skid u kontejneru, spremnik 500 l",
                "ventilacija, izduv, protupožarna zaštita",
                "GRO, DC razvod −48 V, SPD, uzemljenje",
                "preuzimanje opreme Kupca (Azići), montaža FN",
                "NetEco, puštanje u rad, 72 h proba, obuka"])]
    for k, (tag, name, items) in enumerate(lots):
        x0 = 0.4 + k * 4.7
        card(s, x0, 3.1, 4.5, 2.35)
        badge(s, x0 + 0.15, 3.2, 0.62, tag, size=10)
        box(s, x0 + 0.88, 3.2, 3.5, 0.62, [(name, {"bold": True, "size": 11.5})],
            anchor=MSO_ANCHOR.MIDDLE)
        box(s, x0 + 0.2, 3.92, 4.15, 1.85,
            [(t, {"bullet": True, "indent": 0.16}) for t in items], size=11,
            space=3)
    notes(s, "Huawei oprema (ICC360, R4875 ×3, S4875G3 ×2, FN moduli ×12, PVDB, "
             "kablovi) dolazi od Kupca; spisak u Prilogu I, Tačka 9. Ponuđač može "
             "ponuditi jedan ili oba LOT-a.")


def s_postupak(prs, L):
    s = new_slide(prs, L)
    title(s, "Postupak nabavke — od objave do ugovora")
    rows = [("Postupak", "pregovarački postupak sa objavom obavještenja "
                         "(čl. 10 i 16 Pravilnika)"),
            ("Kriterij", "najniža cijena prihvatljive ponude, za svaki LOT"),
            ("Rokovi", "ponude 20 dana od poziva; realizacija 90 dana od uvođenja "
                       "u posao, po LOT-u"),
            ("Plaćanje", "100 %, 30 dana od fakture, po završenom LOT-u i "
                         "primopredaji"),
            ("Garancija", "min. 2 godine + postgarancija 5 godina; kazna 0,1 %/dan, "
                          "max 10 %")]
    y = 1.0
    for k, (lab, val) in enumerate(rows):
        shape(s, MSO_SHAPE.RECTANGLE, 0.4, y, 9.2, 0.4, TINT if k % 2 == 0 else WHITE)
        box(s, 0.5, y, 1.6, 0.4, [(lab, {"bold": True, "size": 10.5, "color": ORANGE})],
            anchor=MSO_ANCHOR.MIDDLE)
        box(s, 2.1, y, 7.4, 0.4, [(val, {"size": 10.5})], anchor=MSO_ANCHOR.MIDDLE,
            space=0)
        y += 0.42
    steps = [("Odluka Uprave", "odobrenje sredstava"),
             ("Zahtjev za nabavku", "TD, komisija, 9 ponuđača"),
             ("Objava i poziv", "obavještenje i poziv"),
             ("Obilazak lokacija", "do 5 dana prije roka"),
             ("Prijem ponuda", "20 dana; jedan ili oba LOT-a"),
             ("Ocjena", "reference, kadrovi, autorizacija DEA"),
             ("Pregovori", "cijena; konačne ponude"),
             ("Izbor i ugovor", "po LOT-u; uvođenje ≤15 dana")]
    w, h, gx = 2.1, 1.12, 0.27
    for k, (head, body) in enumerate(steps):
        r, c = divmod(k, 4)
        x, yy = 0.4 + c * (w + gx), 3.25 + r * (h + 0.18)
        card(s, x, yy, w, h)
        badge(s, x + 0.1, yy + 0.1, 0.38, str(k + 1), size=11)
        box(s, x + 0.55, yy + 0.08, w - 0.62, 0.45, [(head, {"bold": True, "size": 10.5})],
            anchor=MSO_ANCHOR.MIDDLE, space=0)
        box(s, x + 0.12, yy + 0.56, w - 0.22, h - 0.6, [(body, {"size": 9.5})], space=0)
        if c < 3:
            arrow(s, x + w + 0.04, yy + 0.18, 0.19, 0.22)
    notes(s, "Uslovi: min. 2 slična ugovora u 3 godine (LOT 1 ≥10.000 KM, LOT 2 "
             "≥20.000 KM), min. 3 radnika + 1 certificiran serviser, autorizacija "
             "proizvođača DEA, ISO 9001/14001. Primopredaja: Komisija na lokaciji do "
             "15 dana od zahtjeva. Rizik: referentni P18-6 standardno ima SHUNT "
             "pobudu; tražiti AREP+ ili EBS.")


def s_vrijednost(prs, L):
    s = new_slide(prs, L)
    title(s, "Predračunska vrijednost")
    card(s, 0.4, 1.1, 3.4, 2.0)
    box(s, 0.5, 1.2, 3.2, 1.8, [("100.000,00 KM", {"bold": True, "size": 26,
                                                   "color": ORANGE}),
                                ("bez PDV-a, obje lokacije", {"size": 11}),
                                ("sredstva raspoloživa u 2026.", {"size": 11,
                                                                 "color": GREY})],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=2)
    table(s, 4.05, 1.1, 5.55, [["LOT", "Predmet", "Iznos (KM)"],
                               ["LOT 1", "Nosači FN panela, 8 kpl", "30.000,00"],
                               ["LOT 2", "Agregatska postrojenja 18 kVA, 2 kpl", "70.000,00"],
                               ["UKUPNO", "", "100.000,00"]],
          [0.9, 3.15, 1.5], size=11, row_h=0.45, bold_last=True,
          align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT])
    box(s, 0.4, 3.2, 9.2, 0.35, [("Izvor: Trogodišnji plan 2025–2027 i Operativni "
                                  "plan 2026, ID TIRS, tačka 14.6.",
                                  {"size": 10.5, "color": GREY})])
    # dinamika
    box(s, 0.4, 3.62, 4, 0.3, [("Dinamika ulaganja", {"bold": True, "size": 12})])
    box(s, 4.4, 3.65, 5.2, 0.3, [("mjeseci: 10–12/2026 i 1–9/2027; radovi u proljeće/ljeto",
                                  {"size": 9.5, "color": GREY, "italic": True})],
        align=PP_ALIGN.RIGHT)
    months = ["10", "11", "12", "1", "2", "3", "4", "5", "6", "7", "8", "9"]
    acts = [("Odluka Uprave", {0}),
            ("Zahtjev za nabavku", {0, 1}),
            ("Postupak nabavke i ugovor", {1, 2, 3}),
            ("Realizacija (90 dana / LOT)", {5, 6, 7}),
            ("Primopredaja i plaćanje", {8})]
    head = ["Aktivnost"] + months
    rows = [head] + [[a] + [""] * 12 for a, _ in acts]
    tbl = table(s, 0.4, 3.98, 9.2, rows, [3.2] + [0.5] * 12, size=8.5, row_h=0.3,
                align=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 12)
    for i, (_, ms) in enumerate(acts, 1):
        for j in ms:
            c = tbl.cell(i, j + 1)
            c.fill.solid()
            c.fill.fore_color.rgb = ORANGE
    notes(s, "Procjena je iz RFI 'Autonomno napajanje za BS' (27.04.2026) i ranijih "
             "nabavki. Provjeriti LOT 2 prema tržišnoj cijeni agregata sa AREP+/EBS "
             "pobudom prije objave.")


# --------------------------------------------------------------------------
def images(tmp):
    import pymupdf
    out = {}
    d = pymupdf.open(P3)
    pix = pymupdf.Pixmap(d, d[1].get_images()[0][0])
    out["sjednica"] = os.path.join(tmp, "sjednica.jpg")
    Image.frombytes("RGB", (pix.width, pix.height), pymupdf.Pixmap(
        pymupdf.csRGB, pix).samples).save(out["sjednica"], quality=88)
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.Pixmap(d, d[17].get_images()[0][0]))
    out["hamzici"] = os.path.join(tmp, "hamzici.jpg")
    Image.frombytes("RGB", (pix.width, pix.height), pix.samples).rotate(
        -90, expand=True).save(out["hamzici"], quality=88)
    return out


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        JOINT, "TD-OUTPUT", "5. Prezentacija planirane nabavke - hibridni sistemi "
                            "BS Sjednica i BS Hamzići.pptx")
    import tempfile
    tmp = tempfile.mkdtemp()
    img = images(tmp)
    prs, L = prep_template(src)
    edit_title_slide(prs.slides[0])
    s_investicija(prs, L, img)
    s_rjesenje(prs, L)
    s_vrijednost(prs, L)
    s_postupak(prs, L)
    move_last_to_end(prs)
    renumber(prs)
    prs.core_properties.title = "Autonomni hibridni sistemi napajanja BS Sjednica i BS Hamzići"
    prs.save(dst)
    print(f"snimljeno: {dst} ({len(prs.slides)} slajdova)")


if __name__ == "__main__":
    main()
